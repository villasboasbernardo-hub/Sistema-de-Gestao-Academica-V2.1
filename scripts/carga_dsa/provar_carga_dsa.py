"""Prova que a leitura monta os blocos certos e que as recusas da carga do DSA PEGAM.

O QUE  : monta uma planilha SINTETICA do modelo (nomes ficticios), le, e confere bloco a bloco;
         depois arma um retrato de banco em memoria e confere cada recusa pelo caso que a
         discrimina — o instrutor homonimo em outro posto, o vinculo inativo, a UE que falta,
         o 7º TA de TFM.

PARA QUE: validacao que nunca foi vista recusando nao e validacao. Esta carga escreve fora da
         Server Action, entao o que ela nao recusar ninguem recusa: o banco, medido em
         05/10/2026, aceita instrutor nao habilitado.

COMO   : `python -m scripts.carga_dsa.provar_carga_dsa` — saida 0 = todas as provas passam.
         ⚠️ Nao toca em banco nenhum e nao le planilha real.
"""

from __future__ import annotations

import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

from scripts.carga_dsa import planilha, plano, sincronizar
from scripts.carga_dsa.resolver import Pendencia, _aula_em_dia_parado, resolver

FALHAS: list[str] = []


def conferir(nome: str, condicao: bool, detalhe: object = "") -> None:
    print(f"  {'ok ' if condicao else 'FALHOU'}  {nome}" + (f" — {detalhe}" if not condicao else ""))
    if not condicao:
        FALHAS.append(nome)


def planilha_sintetica(pasta: Path) -> Path:
    livro = openpyxl.Workbook()
    catalogo = livro.active
    catalogo.title = planilha.ABA_CATALOGO
    catalogo.append([None])
    catalogo.append(["CÓD", "DISCIPLINA", "Nº UE", "CONCATENADO", "UNIDADES DE ENSINO E TÓPICOS", "CH", "LOCAL", "T/E", "INSTRUTOR/PROFESSOR"])
    for linha in (
        ["I", "DISCIPLINA UM", 1.0, "I1", "1 – TOPICO DA PLANILHA", 3.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", "PE1", "IPE1", "PROVA", 2.0, "SALA 2", "PM", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", "VP", "IVP", "VISTA", 1.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", 9.0, "I9", "9 – UE QUE O BANCO NAO TEM", 1.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["T", "TFM", 1.0, "T1", "1 – CORRIDA", 7.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["AD", "ADMINISTRAÇÃO", 4.0, "AD4", "ESTUDO INDIVIDUAL", None, "LAB INFO", "EO", None],
        ["LP", "LICENÇA", 1.0, "LP1", "LICENÇA DE PAGAMENTO", 7.0, "--", "--", "--"],
        ["I", "REPETIDA", 1.0, "I1", "CHAVE REPETIDA — O PROCV NAO ENXERGA", 99.0, "X", "EO", "X"],
    ):
        catalogo.append(linha)

    folha = livro.create_sheet(planilha.ABA_PREENCHIMENTO)
    folha["A1"], folha["B1"] = "SIGLA DO CURSO:", "C-Exp-QUALQUER"
    folha["A2"], folha["B2"] = "Nº DE ALUNOS:", 6
    folha["B3"], folha["C3"], folha["D3"], folha["E3"] = "DATA", "TA", "LP", 1  # digitado no CABECALHO: nao e TA
    dias = {
        date(2026, 3, 2): {1: ("I", 1.0), 2: ("I", 1.0), 3: ("I", "PE1"), 4: ("I", "PE1"), 6: ("I", 1.0), 8: ("AD", 4.0)},
        date(2026, 3, 3): {n: ("T", 1.0) for n in range(1, 8)},
        date(2026, 3, 4): {1: ("I", "VP"), 2: ("I", 9.0)},
        date(2026, 3, 5): {n: ("LP", 1.0) for n in range(1, 8)},
    }
    r = 4
    for dia, tas in dias.items():
        for n in range(1, 10):
            folha.cell(r, 3, f"{n}º")
            if n == 3:
                folha.cell(r, 2, datetime(dia.year, dia.month, dia.day))
            if n in tas:
                folha.cell(r, 4, tas[n][0])
                folha.cell(r, 5, tas[n][1])
            r += 1
    controle = livro.create_sheet(planilha.ABA_CONTROLE)
    controle.append(["CÓD.", "DISCIPLINA", "CH. PREVISTA", "CH. CUMPRIDA"])
    controle.append(["I", "DISCIPLINA UM", 10.0, 7])  # 3 de aula + 1 na UE 9 + 2 de prova + 1 de vista
    controle.append(["T", "TFM", 40.0, 7])
    caminho = pasta / "sintetica.xlsx"
    livro.save(caminho)
    return caminho


def planilha_da_leva_1(pasta: Path) -> Path:
    """Spec 014 (T107): licenca SO DESTA turma, Monitoria, Recepcao, palestra com responsavel que nao casa,
    e aula no mesmo dia da licenca. Nomes ficticios."""
    livro = openpyxl.Workbook()
    catalogo = livro.active
    catalogo.title = planilha.ABA_CATALOGO
    catalogo.append([None])
    catalogo.append(["CÓD", "DISCIPLINA", "Nº UE", "CONCATENADO", "UNIDADES DE ENSINO E TÓPICOS", "CH", "LOCAL", "T/E", "INSTRUTOR/PROFESSOR"])
    for linha in (
        ["I", "DISCIPLINA UM", 1.0, "I1", "1 – TOPICO DA PLANILHA", 3.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["LA", "LICENÇA", 1.0, "LA1", "LICENÇA ADMINISTRATIVA DA TURMA", None, "--", "--", "--"],
        ["MO", "MONITORIA", 1.0, "MO1", "MONITORIA DE MATEMÁTICA", None, "SALA 2", "--", "CC (T) FULANO SILVA"],
        ["RE", "RECEPÇÃO", 1.0, "RE1", "RECEPÇÃO AOS ALUNOS", None, "SALA 2", "--", "CC (T) FULANO SILVA"],
        ["PA", "PALESTRA", 1.0, "PA1", "PALESTRA DE SEGURANÇA", None, "SALA 2", "--", "SR. CICRANO DE FORA"],
    ):
        catalogo.append(linha)
    folha = livro.create_sheet(planilha.ABA_PREENCHIMENTO)
    folha["A1"], folha["B1"] = "SIGLA DO CURSO:", "C-Exp-QUALQUER"
    folha["A2"], folha["B2"] = "Nº DE ALUNOS:", 6
    dias = {
        date(2026, 3, 9): {1: ("LA", 1.0), 2: ("LA", 1.0), 3: ("LA", 1.0), 4: ("LA", 1.0), 6: ("I", 1.0)},
        date(2026, 3, 10): {1: ("MO", 1.0), 2: ("RE", 1.0), 3: ("PA", 1.0), 4: ("I", 1.0), 5: ("I", 1.0)},
    }
    r = 4
    for dia, tas in dias.items():
        for n in range(1, 10):
            folha.cell(r, 3, f"{n}º")
            if n == 3:
                folha.cell(r, 2, datetime(dia.year, dia.month, dia.day))
            if n in tas:
                folha.cell(r, 4, tas[n][0])
                folha.cell(r, 5, tas[n][1])
            r += 1
    controle = livro.create_sheet(planilha.ABA_CONTROLE)
    controle.append(["CÓD.", "DISCIPLINA", "CH. PREVISTA", "CH. CUMPRIDA"])
    controle.append(["I", "DISCIPLINA UM", 12.0, 3])  # o curriculo diz 10: erro de planilha, sem mudar o catalogo
    caminho = pasta / "leva1.xlsx"
    livro.save(caminho)
    return caminho


def planilha_da_leva_2(pasta: Path) -> Path:
    """Spec 014, PR 6 (T223): duas provas mistas e uma de recuperacao, a vista da PM 1 em DUAS sessoes, o fiscal e o
    2º instrutor. Nomes ficticios."""
    livro = openpyxl.Workbook()
    catalogo = livro.active
    catalogo.title = planilha.ABA_CATALOGO
    catalogo.append([None])
    catalogo.append(["CÓD", "DISCIPLINA", "Nº UE", "CONCATENADO", "UNIDADES DE ENSINO E TÓPICOS", "CH", "LOCAL", "T/E", "INSTRUTOR/PROFESSOR"])
    for linha in (
        ["I", "DISCIPLINA UM", 1.0, "I1", "1 – TOPICO DA PLANILHA", 3.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", "PM1", "IPM1", "PROVA MISTA 1", 3.0, "SALA 2", "PM", "CC (T) FULANO SILVA / SR. AUXILIAR DE FORA"],
        ["I", "DISCIPLINA UM", "PM2", "IPM2", "PROVA MISTA 2", 2.0, "SALA 2", "PM", "CC (T) FULANO SILVA (FISCAL)"],
        ["I", "DISCIPLINA UM", "RECON", "IRECON", "PROVA DE RECUPERACAO", 1.0, "SALA 2", "PM", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", "VP1", "IVP1", "VISTA DA PROVA", 1.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
        ["I", "DISCIPLINA UM", "PP1", "IPP1", "PROVA PRÁTICA", 1.0, "SALA 2", "PP", "CC (T) FULANO SILVA"],
    ):
        catalogo.append(linha)
    folha = livro.create_sheet(planilha.ABA_PREENCHIMENTO)
    folha["A1"], folha["B1"] = "SIGLA DO CURSO:", "C-Exp-QUALQUER"
    folha["A2"], folha["B2"] = "Nº DE ALUNOS:", 6
    dias = {
        date(2026, 4, 6): {1: ("I", 1.0), 2: ("I", 1.0), 3: ("I", 1.0), 4: ("I", "PM1"), 5: ("I", "PM1"), 6: ("I", "PM1")},
        date(2026, 4, 7): {1: ("I", "VP1")},
        date(2026, 4, 8): {1: ("I", "VP1"), 3: ("I", "PM2"), 4: ("I", "PM2")},
        date(2026, 4, 9): {1: ("I", "RECON"), 3: ("I", "PP1")},
    }
    r = 4
    for dia, tas in dias.items():
        for n in range(1, 10):
            folha.cell(r, 3, f"{n}º")
            if n == 3:
                folha.cell(r, 2, datetime(dia.year, dia.month, dia.day))
            if n in tas:
                folha.cell(r, 4, tas[n][0])
                folha.cell(r, 5, tas[n][1])
            r += 1
    controle = livro.create_sheet(planilha.ABA_CONTROLE)
    controle.append(["CÓD.", "DISCIPLINA", "CH. PREVISTA", "CH. CUMPRIDA"])
    controle.append(["I", "DISCIPLINA UM", 10.0, 12])  # 3 de aula + 3 PM1 + 2 vistas + 2 PM2 + 1 RECON + 1 PP1
    caminho = pasta / "leva2.xlsx"
    livro.save(caminho)
    return caminho


ITENS_DA_LEVA_2 = [
    {"id": "av-1", "disciplina_id": "d-1", "numero": 1, "instrumento": "PM", "instrumento_nome": "Prova Mista",
     "ta_aplicacao": 2, "origem_ta_aplicacao": "curriculo", "ta_vista": 1, "origem_ta_vista": "curriculo"},
    {"id": "av-2", "disciplina_id": "d-1", "numero": 2, "instrumento": "PM", "instrumento_nome": "Prova Mista",
     "ta_aplicacao": None, "origem_ta_aplicacao": None, "ta_vista": None, "origem_ta_vista": None},
]


# A marca de vista obrigatoria do banco (migration D-9, 16/10/2026: PM, PO e PE) — DADO DE TESTE, espelho do banco.
COM_VISTA = frozenset({"PM", "PO", "PE"})


def retrato(habilitado: str = "ativo") -> dict:
    """O banco em memoria. Dois «FULANO SILVA» em postos diferentes: o caso que discrimina."""
    return {
        "turma": {"id": "t-1", "codigo": "TURMA SINTETICA 2026", "curso_id": "c-1", "curso": "C-Exp-QUALQUER"},
        "autor": {"codigo": "USR-X", "auth_user_id": "a-1", "status": "ativo"},
        "disciplinas": [
            {"id": "d-1", "cod": "I", "nome": "DISCIPLINA UM", "ch": 10},
            {"id": "d-t", "cod": "T", "nome": "Treinamento Físico Militar", "ch": 40},
        ],
        "ues": [
            {"id": "u-1", "codigo": "UE-1", "disciplina_id": "d-1", "numero": 1, "topico": "TOPICO DO CATALOGO DO BANCO", "ch": 3},
            {"id": "u-t", "codigo": "UE-T", "disciplina_id": "d-t", "numero": 1, "topico": "CORRIDA", "ch": 40},
        ],
        "instrutores": [
            {"id": "i-cc", "codigo": "1", "posto": "CC", "esp": "(T)", "nome": "FULANO DE TAL SILVA SOUZA", "guerra": None},
            {"id": "i-sg", "codigo": "2", "posto": "3ºSG", "esp": "-GC", "nome": "FULANO SILVA PEREIRA", "guerra": None},
        ],
        "habilitacoes": [
            {"instrutor_id": "i-cc", "disciplina_id": "d-1", "status": habilitado},
            {"instrutor_id": "i-cc", "disciplina_id": "d-t", "status": "ativo"},
        ],
        "metodologias": [{"valor": "Exposição Oral", "sigla": "EO"}, {"valor": "Prova Mista", "sigla": "PM"}],
        "tipos_atividade": [{"valor": "Estudo Individual", "categoria": "Estudo_Individual"}],
        "tipos_avaliacao": ["Prova Escrita"],
        "salas": [],
        "tetos": {"dsa.teto_tfm_semana": "6", "dsa.teto_recomendado_semana": "25"},
        "feriados": [],
        "da_turma": [],
        "codigos": {"registros_aula": [], "avaliacoes": [], "atividades_nao_letivas": []},
        "de_outras_turmas": [],
        # Emenda de 15/10/2026 (D-9, 16/10/2026): a janela e a marca de vista sao DADO do banco — aqui, espelho de teste.
        "janela_sessoes_dias": "7",
        "instrumentos_com_vista": sorted(COM_VISTA),
        "instrumentos": [{"valor": v, "nome": n} for v, n in (("PE", "Prova Escrita"), ("PM", "Prova Mista"), ("PP", "Prova Prática"),
                                                               ("PO", "Prova Oral"), ("TG", "Trabalho em Grupo"),
                                                               ("TI", "Trabalho Individual"), ("OD", "Observação de Desempenho"))],
    }


DECISOES = {
    "turma": "TURMA SINTETICA 2026",
    "locais": {"SALA 2": "Sala 02", "LAB INFO": "Laboratório de Informática"},
    "tipo_avaliacao": "Prova Escrita",
    "padrao_ue_de_avaliacao": "PE[0-9]+",
    "padrao_ue_de_vista": "VP",
    "chaves": {"AD+4": {"tipo": "estudo_individual"}, "LP+1": {"tipo": "calendario", "descricao": "Licença"}},
    "vistas": [{"disciplina": "I", "data": "2026-03-04", "da_prova_de": "2026-03-02"}],
}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    with tempfile.TemporaryDirectory() as pasta:
        leitura = planilha.ler(planilha_sintetica(Path(pasta)))

    print("LEITURA")
    blocos = [(b.data.isoformat(), b.ta_inicial, b.tempos, b.cod, b.ue) for b in leitura.blocos]
    conferir("TA consecutivos com a mesma chave viram UM bloco", ("2026-03-02", 1, 2, "I", "1") in blocos, blocos)
    conferir("buraco no meio abre bloco novo", ("2026-03-02", 6, 1, "I", "1") in blocos, blocos)
    conferir("chave diferente abre bloco novo", ("2026-03-02", 3, 2, "I", "PE1") in blocos, blocos)
    conferir("o bloco guarda a LINHA em que comeca", leitura.blocos[0].linha == 4, leitura.blocos[0].linha)
    conferir("digitado no cabecalho NAO e TA e volta em `ignoradas`",
             [(c.celula, c.valor) for c in leitura.ignoradas] == [("D3", "LP"), ("E3", "1")], leitura.ignoradas)
    conferir("chave repetida no catalogo: vale a primeira, como no PROCV",
             leitura.catalogo[("I", "1")].ch == 3 and len(leitura.duplicadas) == 1, leitura.duplicadas)

    print("INSTRUTOR — posto + nome de guerra, nunca pedaco do nome")
    lido = plano.ler_instrutor("CC (T) FULANO SILVA (FISCAL)")
    achados = plano.casar_instrutor(lido, retrato()["instrutores"])
    conferir("dois «FULANO SILVA»: o posto decide, e casa exatamente um", [a["codigo"] for a in achados] == ["1"], achados)
    sem_posto = [i for i in retrato()["instrutores"] if all(p in plano.normalizar(i["nome"]).split() for p in lido.guerra.split())]
    conferir("controle: SEM o posto os dois casariam (o teste discrimina)", len(sem_posto) == 2, sem_posto)
    conferir("«SO-ME (RM1) BELTRANO» → posto SO, especialidade ME RM1",
             (plano.ler_instrutor("SO-ME (RM1) BELTRANO").posto, plano.ler_instrutor("SO-ME (RM1) BELTRANO").especialidade) == ("SO", "ME RM1"))

    print("RN-INST-01 — habilitacao")
    conferir("ministrar sem vinculo: recusa", not plano.pode_atuar("ministrar", "i-sg", "d-1", retrato()["habilitacoes"]))
    conferir("ministrar com vinculo INATIVO: recusa", not plano.pode_atuar("ministrar", "i-cc", "d-1", retrato("inativo")["habilitacoes"]))
    conferir("avaliacao NAO exige habilitacao", plano.pode_atuar("avaliacao", "i-sg", "d-1", []))

    print("O PLANO")
    p = plano.montar(leitura, retrato(), DECISOES, "sintetica.xlsx", "2026-03-10")
    recusas = "\n".join(p.recusas)
    conferir("UE que o banco nao tem: recusa, e nao cria", "UE 9 da disciplina I NAO existe" in recusas, recusas)
    conferir("7 TA de TFM numa semana: recusa (RN-DIST-03)", "teto de TFM" in recusas, recusas)
    conferir("UE × CH do catalogo da planilha: a UE 9, que nao entrou, nao fecha — recusa",
             "UE I+9: lancado 0 TA" in recusas, recusas)
    conferir("disciplina × CONTROLE: I soma 6 no plano e a planilha conta 7 — recusa",
             "disciplina I: o plano soma 6 TA" in recusas, recusas)
    conferir("controle positivo: T fecha com a UE (7) e com o CONTROLE (7), e nao e acusada",
             "UE T+1" not in recusas and "disciplina T:" not in recusas, recusas)
    conferir("so essas quatro recusas", len(p.recusas) == 4, p.recusas)
    conferir("o titulo da prova vem da planilha", p.avaliacoes[0]["conteudo_resumo"] == "PROVA", p.avaliacoes[0])
    andamento = plano.montar(leitura, retrato(), {**DECISOES, "turma_em_andamento": True}, "s.xlsx", "2026-03-10")
    conferir("turma em andamento: UE abaixo da CH vira alerta, nao recusa",
             not any(r.startswith("UE I+9") for r in andamento.recusas) and any("UE I+9" in x for x in andamento.alertas))
    conferir("o topico gravado e o do BANCO, nunca o da planilha",
             {a["conteudo_resumo"] for a in p.aulas if a["disciplina"] == "I"} == {"TOPICO DO CATALOGO DO BANCO"})
    conferir("a vista vai na MESMA linha da avaliacao (RN-AVAL-02)",
             len(p.avaliacoes) == 1 and p.avaliacoes[0]["data_vista_prova"] == "2026-03-04", p.avaliacoes)
    conferir("dia de calendario NAO vira lancamento", [c["data"] for c in p.calendario] == ["2026-03-05"] and
             not any(a["data"] == "2026-03-05" for a in p.atividades))
    conferir("codigo deterministico, derivado da linha", p.aulas[0]["codigo"] == "DSAP-TURMA-SINTETICA-2026-L004-A", p.aulas[0]["codigo"])

    p_inativo = plano.montar(leitura, retrato("inativo"), DECISOES, "sintetica.xlsx", "2026-03-10")
    conferir("instrutor com vinculo inativo: as aulas dele sao recusadas",
             sum("nao esta habilitado" in r for r in p_inativo.recusas) == 2, p_inativo.recusas)

    seis = retrato()
    seis["tetos"]["dsa.teto_tfm_semana"] = "7"
    conferir("o teto vem do PARAMETRO: com 7 no banco, os mesmos 7 TA passam",
             not any("teto de TFM" in r for r in plano.montar(leitura, seis, DECISOES, "s.xlsx", "2026-03-10").recusas))

    print("A TRANSACAO")
    ensaio, gravar = plano.sql(p, DECISOES, ensaio=True), plano.sql(p, DECISOES, ensaio=False)
    conferir("o ensaio termina levantando a sentinela", plano.SENTINELA_DO_ENSAIO in ensaio and "raise exception" in ensaio)
    conferir("a gravacao NAO levanta a sentinela", plano.SENTINELA_DO_ENSAIO not in gravar)
    conferir("nenhuma linha nasce com procedencia de ETL (seria «herdada» na grade)", "origem_migracao_v1" not in gravar)
    conferir("toda insercao e idempotente pelo codigo", gravar.count("on conflict (codigo) do nothing") == 3, gravar.count("on conflict"))

    print("SPEC 014 — LEVA 1 (resolver da sincronizacao)")
    with tempfile.TemporaryDirectory() as pasta:
        leva1 = planilha.ler(planilha_da_leva_1(Path(pasta)))
    ref = retrato()
    ref["tipos_atividade"] = [
        {"valor": "Estudo Individual", "categoria": "Estudo_Individual"},
        {"valor": "Monitoria", "categoria": "Apoio_Individualizado"},
        {"valor": "Recepção aos alunos", "categoria": "TAD"},
        {"valor": "Palestra", "categoria": "AEC"},
    ]
    dec = {"turma": "TURMA SINTETICA 2026", "locais": {"SALA 2": "Sala 02"}, "chaves": {}}
    res = resolver(leva1, ref, dec, "leva1.xlsx", set(), False)
    conferir("FR-024 · licenca SO DESTA turma vira dia parado previsto, tipo LA, codigo …-P",
             [(d["data"], d["ta_inicial"], d["tempos"], d["tipo_parada"], d["prevista"], d["codigo"][-6:]) for d in res.dias_parados]
             == [("2026-03-09", 1, 4, "LA", True, "L004-P")], res.dias_parados)
    conferir("FR-024 · e NAO vira TAD (era o que inflava o TAD, D-AA-12)",
             not any(a["data"] == "2026-03-09" and a["categoria_normativa"] == "TAD" for a in res.atividades), res.atividades)
    conferir("FR-024 · dia GLOBAL continua calendario (feriados), nao dia parado",
             [c["data"] for c in resolver(leva1, ref, dec, "l.xlsx", {"2026-03-09"}, True).calendario] == ["2026-03-09"])
    cat = {a["subtipo"]: a["categoria_normativa"] for a in res.atividades}
    conferir("FR-021 · MONITORIA vira Apoio Individualizado", cat.get("Monitoria") == "Apoio_Individualizado", cat)
    conferir("FR-021 · RECEPCAO e TAD com o tipo «Recepção aos alunos»", cat.get("Recepção aos alunos") == "TAD", cat)
    # RN-RESP-1 (emenda de 15/10/2026): a palestra aceita responsavel de nome livre. O VEREDITO VIRA (DoD 8): ate
    # 15/10 o texto fora do cadastro virava `responsavel_sem_decisao` e a palestra ficava fora (FR-016, A3.2).
    palestra = [a for a in res.atividades if a["subtipo"] == "Palestra"]
    conferir("RN-RESP-1 · palestra com responsavel fora do cadastro ENTRA, com o texto como responsavel livre, sem pendencia",
             len(palestra) == 1 and palestra[0]["instrutor_id"] is None and "CICRANO" in (palestra[0]["responsavel_externo"] or "")
             and not any(pp.tipo == "responsavel_sem_decisao" for pp in res.pendencias), (palestra, [pp.detalhe for pp in res.pendencias]))
    conferir("RN-RESP-1 · e a linha fica registrada para a lista de conferencia (fora do git)",
             any("CICRANO" in x["texto"] and x["tipo"] == "aec" for x in res.responsavel_livre), res.responsavel_livre)
    conferir("RN-RESP-1 · a aula continua exigindo instrutor do cadastro (o plano tem aula, e nenhuma sem instrutor)",
             bool(res.aulas) and all(a["instrutor_id"] for a in res.aulas), res.aulas)
    # Duvida 6 (Bernardo Villas Boas, 16/10/2026): o texto digitado por cima da formula vale tambem para a ATIVIDADE.
    import copy as _copy
    digitada = _copy.deepcopy(leva1)
    linha_pa = next(b.linha for b in digitada.blocos if b.cod == "PA")
    digitada.digitadas[linha_pa] = {"topico": "PALESTRA DIGITADA NA IMPRESSAO"}
    com_digitado = [a["descricao"] for a in resolver(digitada, ref, dec, "l.xlsx", set(), True).atividades if a["subtipo"] == "Palestra"]
    sem_digitado = [a["descricao"] for a in res.atividades if a["subtipo"] == "Palestra"]
    conferir("duvida 6 · DISCRIMINA: o topico digitado na IMPRESSAO vira a descricao da atividade (sem ele, o do catalogo)",
             com_digitado == ["PALESTRA DIGITADA NA IMPRESSAO"] and sem_digitado == ["PALESTRA DE SEGURANÇA"], (com_digitado, sem_digitado))
    from scripts.carga_dsa.resolver import vazio
    conferir("RN-RESP-1 · celula so de tracos («------») e vazia, nao nome", vazio("------") and vazio(" -- ") and not vazio("CT FULANO"))
    # A chave `responsavel_externo_aceito` do lote ficou SEM EFEITO com a RN-RESP-1: o texto entra igual com ou sem ela.
    aceito = resolver(leva1, ref, {**dec, "responsavel_externo_aceito": ["SR. CICRANO DE FORA"]}, "l.xlsx", set(), True)
    sem_lote = resolver(leva1, ref, dec, "l.xlsx", set(), True)
    conferir("RN-RESP-1 · a decisao do lote (`responsavel_externo_aceito`) nao muda nada: com ou sem ela, o mesmo plano",
             aceito.impressao_digital() == sem_lote.impressao_digital(), (aceito.impressao_digital(), sem_lote.impressao_digital()))
    # R5 (decisao de Bernardo Villas Boas, 16/10/2026; emenda da FR-036.1 e da D3.6): a conferencia e por TEMPO. A licenca
    # de 09/03 ocupa os tempos 1–4 e a aula da I, o 6º: parada de PARTE do dia, sem conflito — por dia ela acusava.
    conferir("R5 · DISCRIMINA: parada de parte do dia e aula em OUTRO tempo — a aula entra e NAO vira erro (por dia, virava)",
             not any("marca parado" in e["diverge"] for e in res.erros_tabelados)
             and any(a["data"] == "2026-03-09" for a in res.aulas), res.erros_tabelados)
    sobreposto = type(res)(turma=res.turma, autor=res.autor, titulo="x")
    sobreposto.aulas.append({"data": "2026-03-09", "ta_inicial": 3, "tempos_consumidos": 2, "disciplina": "I", "ue": "UE-1"})
    sobreposto.dias_parados.append({"data": "2026-03-09", "ta_inicial": 4, "tempos": 2, "tipo_parada": "LA", "motivo": "LICENCA"})
    _aula_em_dia_parado(sobreposto)
    conferir("R5 · controle: aula num TEMPO parado (3–4 × 4–5) continua erro de planilha, e a aula entra",
             any("marca parado" in e["diverge"] for e in sobreposto.erros_tabelados), sobreposto.erros_tabelados)
    conferir("§11-B · CH da disciplina: planilha 12, curriculo 10 — erro de planilha",
             any(e["diverge"] == "CH da disciplina" and e["planilha"].startswith("12") and e["curriculo"].startswith("10")
                 for e in res.erros_tabelados), res.erros_tabelados)
    conferir("o plano e deterministico (duas resolucoes, a mesma impressao digital)",
             res.impressao_digital() == resolver(leva1, ref, dec, "leva1.xlsx", set(), False).impressao_digital())
    sql = sincronizar.sql_da_turma(res, True, "Nacional", ensaio=False)
    conferir("o dia parado e gravado em `dias_parados`, com `on conflict (codigo)`",
             "insert into public.dias_parados" in sql and sql.count("on conflict (codigo)") >= 2)
    conferir("a linha TAD antiga da mesma linha (…-N) sai com a marca de reclassificacao",
             sincronizar.MARCA_DE_RECLASSIFICACAO in sql and "L004-N" in sql)
    conferir("feriado pelo DEFAULT `FER-`, nunca `max + 1` (E-2)", "max(substring(codigo from '^FER-" not in sql)
    dif = sincronizar.diferencas(res, {"aulas": [], "avaliacoes": [], "atividades": [], "dias_parados": []})
    conferir("diferencas conhece `dias_parados`", bool(dif.get("dias_parados", {}).get("novas")), dif.get("dias_parados"))
    vazia = {c: None for c in sincronizar.CAMPOS["avaliacoes"]}
    falsa = type(res)(turma=res.turma, autor=res.autor, titulo="x")
    falsa.avaliacoes.append({**vazia, "codigo": "X"})
    conferir("FR-049 · `cancelada_na_tela` nao conta como reativada (zero mudanca na 2ª rodada)",
             not sincronizar.diferencas(falsa, {"aulas": [], "avaliacoes": [{**vazia, "codigo": "X", "status": "cancelada_na_tela"}],
                                                "atividades": [], "dias_parados": []})["avaliacoes"]["reativadas"])
    conferir("controle: a cancelada pela CARGA (inativo) volta como reativada",
             sincronizar.diferencas(falsa, {"aulas": [], "avaliacoes": [{**vazia, "codigo": "X", "status": "inativo"}],
                                            "atividades": [], "dias_parados": []})["avaliacoes"]["reativadas"] == ["X"])

    print("A SISTEMA VENCE (A3.3 virada por turma, A3.4 edicao da tela vence)")
    vazio_aula = {c: None for c in sincronizar.CAMPOS["aulas"]}

    def plano_de(*aulas):
        r = type(res)(turma=res.turma, autor=res.autor, titulo="x")
        r.aulas.extend({**vazio_aula, "observacoes": "", "unidade_ensino_id": None, "disciplina_id": None, "instrutor_id": None, **a} for a in aulas)
        return r

    def ref_com(*atuais, de_tela=()):
        return {"atuais": {"aulas": [{**vazio_aula, "status": "ativo", "na_tela": False, **a} for a in atuais],
                           "avaliacoes": [], "atividades": [], "dias_parados": []},
                "de_tela_linhas": list(de_tela)}

    r = plano_de({"codigo": "A", "data": "2026-10-09", "ta_inicial": 1, "tempos_consumidos": 2},
              {"codigo": "B", "data": "2026-10-13", "ta_inicial": 1, "tempos_consumidos": 2})
    rf = ref_com({"codigo": "C", "data": "2026-10-14", "ta_inicial": 3, "tempos_consumidos": 1})
    prot, _d = sincronizar.proteger(r, rf, "2026-10-13")
    conferir("A3.3 · a linha da planilha NA data da virada sai do plano; a de antes fica",
             [a["codigo"] for a in r.aulas] == ["A"], r.aulas)
    conferir("A3.3 · a linha do banco depois da virada fica protegida e nao «sai»",
             "C" in prot["aulas"] and sincronizar.diferencas(r, rf["atuais"], prot)["aulas"]["saem"] == [], prot)
    conferir("controle: SEM virada, a mesma linha do banco «sai»",
             sincronizar.diferencas(r, rf["atuais"])["aulas"]["saem"] == ["C"])
    sql = sincronizar.sql_da_turma(r, False, None, ensaio=False, virada="2026-10-13", protegidas=prot)
    conferir("A3.3 · o SQL so inativa ETL e linha que saiu ANTES da virada, e nunca a protegida",
             sql.count("< '2026-10-13'::date") >= 4 and "array['C']::text[]" in sql)
    # Achado do ensaio de 16/10/2026: a vista 2 lancada na tela para 23/10, numa prova de 06/08 da CAHO, era inativada —
    # o filtro olhava a data da PROVA. A sessao (de vista e de aplicacao) da virada em diante e do sistema.
    conferir("A3.3 · a SESSAO da virada em diante fica: as duas inativacoes de sessao filtram a data DA SESSAO",
             sql.count("s.data < '2026-10-13'::date") == 2, sql.count("s.data < '2026-10-13'::date"))

    r = plano_de({"codigo": "E", "data": "2026-10-05", "ta_inicial": 1, "tempos_consumidos": 5, "local": "Sala 01"})
    rf = ref_com({"codigo": "E", "data": "2026-10-05", "ta_inicial": 1, "tempos_consumidos": 4, "local": "EAD", "na_tela": True})
    prot, div = sincronizar.proteger(r, rf, None)
    conferir("A3.4 · a linha editada na tela sai do plano (nao e desfeita) e vai para «planilha diverge do sistema»",
             r.aulas == [] and "E" in prot["aulas"] and len(div) == 1 and "tempos_consumidos" in div[0]["campos"], div)
    dif = sincronizar.diferencas(r, rf["atuais"], prot)["aulas"]
    conferir("A3.4 · e a rodada da zero mudanca sobre ela", not any(dif.values()), dif)
    r = plano_de({"codigo": "E", "data": "2026-10-05", "ta_inicial": 1, "tempos_consumidos": 5, "local": "Sala 01"})
    rf2 = ref_com({"codigo": "E", "data": "2026-10-05", "ta_inicial": 1, "tempos_consumidos": 4, "local": "EAD", "na_tela": False})
    sincronizar.proteger(r, rf2, None)
    conferir("controle: a MESMA linha sem edicao da tela e atualizada pela planilha",
             sincronizar.diferencas(r, rf2["atuais"])["aulas"]["mudadas"] != [])

    r = plano_de({"codigo": "F", "data": "2026-10-06", "ta_inicial": 3, "tempos_consumidos": 2},
              {"codigo": "G", "data": "2026-10-06", "ta_inicial": 5, "tempos_consumidos": 2})
    prot, div = sincronizar.proteger(r, ref_com(de_tela=[{"tabela": "aulas", "codigo": "DSA-X", "data": "2026-10-06", "ta_inicial": 4, "tempos": 1}]), None)
    conferir("A3.4 · o lancamento NASCIDO na tela fica; a linha da planilha no mesmo tempo sai e e listada, a vizinha entra",
             [a["codigo"] for a in r.aulas] == ["G"] and div and "DSA-X" in div[0]["motivo"], (r.aulas, div))

    r = plano_de()
    r.pendentes = {"P-N"}
    vazio_ativ = {c: None for c in sincronizar.CAMPOS["atividades"]}
    rf = {"atuais": {"aulas": [], "avaliacoes": [], "dias_parados": [],
                     "atividades": [{**vazio_ativ, "codigo": "P-N", "data": "2026-10-05", "status": "ativo", "na_tela": False},
                                    {**vazio_ativ, "codigo": "Q-N", "data": "2026-10-05", "status": "ativo", "na_tela": False}]},
          "de_tela_linhas": []}
    prot, _d = sincronizar.proteger(r, rf, None)
    conferir("A3.2 · a linha que ficou fora por pendencia de responsavel e JA esta no banco nao «sai»; a que saiu da planilha sai",
             sincronizar.diferencas(r, rf["atuais"], prot)["atividades"]["saem"] == ["Q-N"], prot)

    conferir("A3.2 · proposta de classificacao do responsavel (entidade · pessoa · palestrante)",
             sincronizar.classificar_responsavel("CIAARA-10", set()) == "entidade"
             and sincronizar.classificar_responsavel("1ºTEN FULANO", {"FULANO"}).startswith("pessoa cadastrada")
             and sincronizar.classificar_responsavel("SR. CICRANO DE FORA", set()) == "palestrante externo")

    print("A3.3 — a pendencia de bloco da virada em diante nao para a turma (decisao de 26/10/2026, C-Ap-HN)")
    sem_vinculo = {**retrato("inativo"), "tipos_atividade": ref["tipos_atividade"]}
    hab = [p for p in resolver(leva1, sem_vinculo, dec, "l.xlsx", set(), False).pendencias if p.tipo == "sem_habilitacao"]
    conferir("a pendencia de BLOCO guarda o dia do bloco (as duas aulas de I, 09/03 e 10/03, sem habilitacao)",
             sorted(p.data for p in hab) == ["2026-03-09", "2026-03-10"], [(p.tipo, p.data) for p in hab])
    param, depois = sincronizar.pendencias_que_bloqueiam(hab, set(), "2026-03-10")
    conferir("com virada em 10/03: a de 09/03 PARA a turma; a de 10/03 e listada como depois da virada",
             [p.data for p in param] == ["2026-03-09"] and [p.data for p in depois] == ["2026-03-10"], (param, depois))
    conferir("controle: SEM virada as duas param (era o caso da C-Ap-HN antes de 26/10/2026 — o teste discrimina)",
             len(sincronizar.pendencias_que_bloqueiam(hab, set(), None)[0]) == 2)
    da_turma = Pendencia("autor", "T", "a conta autora nao existe", 0, True, "")
    conferir("a pendencia SEM dia (da turma inteira) para mesmo com virada; a adiada por decisao nao para",
             sincronizar.pendencias_que_bloqueiam([da_turma], set(), "2026-03-10")[0] == [da_turma]
             and sincronizar.pendencias_que_bloqueiam(hab, {"sem_habilitacao"}, None)[0] == [])

    print("SPEC 014 — LEVA 2 (o item, as sessoes, os papeis e as sugestoes)")
    with tempfile.TemporaryDirectory() as pasta:
        leva2 = planilha.ler(planilha_da_leva_2(Path(pasta)))
    ref2 = retrato()
    ref2["tipos_avaliacao"] = ["Prova Escrita", "Prova Mista", "Prova Prática"]
    ref2["itens_curriculo"] = ITENS_DA_LEVA_2
    ref2["ajustes"] = []
    dec2 = {"turma": "TURMA SINTETICA 2026", "locais": {"SALA 2": "Sala 02"}, "chaves": {},
            "padrao_ue_de_avaliacao": "(PM|PP)[0-9]+|RECON", "padrao_ue_de_vista": "VP[0-9]*",
            "tipo_avaliacao_por_tecnica": {"PM": "Prova Mista", "PP": "Prova Prática"}, "vista_extra_como_aec": True}
    res2 = resolver(leva2, ref2, dec2, "leva2.xlsx", set(), True)
    prova = {a["chave_ue"]: a for a in res2.avaliacoes}
    conferir("FR-039 · a PM 1 da planilha aponta o item do curriculo (instrumento e numero)",
             prova.get("PM1", {}).get("avaliacao_curriculo_id") == "av-1", res2.avaliacoes)
    conferir("FR-039 · a PM 2 aponta o item 2", prova.get("PM2", {}).get("avaliacao_curriculo_id") == "av-2")
    conferir("FR-036.1 · TA da PM 1 (3) diferente do curriculo (2): erro de planilha, e a prova entra com o da planilha",
             any("TA de prova diferente" in e and "PM 1" in e for e in res2.erros_de_planilha)
             and prova["PM1"]["tempos_consumidos"] == 3, res2.erros_de_planilha)
    conferir("FR-036.1 · a linha RECON entra SEM item e vai para a lista",
             prova.get("RECON", {}).get("avaliacao_curriculo_id") is None
             and any("RECON" in e for e in res2.erros_de_planilha), res2.erros_de_planilha)
    conferir("FR-036 · a PROVA PRÁTICA numa disciplina que so preve PM NAO se liga a nada (sem inferencia) e vai para a lista",
             prova.get("PP1", {}).get("avaliacao_curriculo_id") is None
             and any("PP1" in e and "nao casa" in e for e in res2.erros_de_planilha), res2.erros_de_planilha)
    from scripts.carga_dsa.resolver import Resolvido, _ligar_itens
    ordem = Resolvido(turma={}, autor={}, titulo="x")
    base = {"disciplina": "I", "disciplina_id": "d-1", "chave_ue": "AV", "tipo_avaliacao": "Prova Escrita",
            "tempos_consumidos": 2, "ta_inicial": 1, "avaliacao_curriculo_id": None, "item": None}
    ordem.avaliacoes = [{**base, "data_avaliacao": "2026-04-01", "descricao": "PM2 – ESTATISTICA (U.E. 2)"},
                        {**base, "data_avaliacao": "2026-04-20", "descricao": "PM1 – MATEMATICA (U.E. 1)"}]
    _ligar_itens(ordem, {"itens_curriculo": ITENS_DA_LEVA_2})
    conferir("FR-036 · o numero ESCRITO na descricao («PM2 –») vence a ordem de aplicacao: a PM 2 aplicada antes liga ao item 2",
             [a["avaliacao_curriculo_id"] for a in ordem.avaliacoes] == ["av-2", "av-1"], ordem.avaliacoes)
    # 14/10/2026 (Bernardo Villas Boas): a prova sem item nao some da lista — a da disciplina sem avaliacao do
    # curriculo carregada, e a vista nao lancada de prova sem item (que o alerta do sistema deixou de contar).
    from datetime import date as _date
    from scripts.carga_dsa.resolver import _vista_de_prova_sem_item
    sem = Resolvido(turma={}, autor={}, titulo="x")
    sem.avaliacoes = [{**base, "disciplina": "IX", "disciplina_id": "d-9", "data_avaliacao": "2026-04-01",
                       "descricao": "PROVA ESCRITA", "sessoes": []}]
    _ligar_itens(sem, {"itens_curriculo": ITENS_DA_LEVA_2})
    conferir("A4.7 · a prova da disciplina SEM avaliacao do curriculo carregada vai para a lista (antes, sumia)",
             any("IX AV" in e and "nao tem avaliacao do curriculo carregada" in e for e in sem.erros_de_planilha),
             sem.erros_de_planilha)
    _vista_de_prova_sem_item(sem, _date(2026, 4, 9), COM_VISTA)
    conferir("vista · a prova sem item, sem vista, no 8º dia vai para a lista", any("vista nao lancada" in e for e in sem.erros_de_planilha),
             sem.erros_de_planilha)
    cedo = Resolvido(turma={}, autor={}, titulo="x")
    cedo.avaliacoes = [dict(sem.avaliacoes[0])]
    _vista_de_prova_sem_item(cedo, _date(2026, 4, 8), COM_VISTA)
    ligada = Resolvido(turma={}, autor={}, titulo="x")
    ligada.avaliacoes = [{**sem.avaliacoes[0], "avaliacao_curriculo_id": "av-1"}]
    _vista_de_prova_sem_item(ligada, _date(2026, 4, 30), COM_VISTA)
    conferir("controle: no 7º dia, ou com item ligado (o alerta do sistema cobre), nao vai para a lista",
             cedo.erros_de_planilha == [] and ligada.erros_de_planilha == [], (cedo.erros_de_planilha, ligada.erros_de_planilha))
    sessoes = prova["PM1"]["sessoes"]
    conferir("FR-042 · a vista da PM 1 em DUAS sessoes, numeradas 1 e 2, cada uma no seu dia",
             [(v["numero"], v["data"]) for v in sessoes] == [(1, "2026-04-07"), (2, "2026-04-08")], sessoes)
    conferir("FR-042 · e nenhuma AEC «vista_extra_como_aec» (a chave fica sem efeito)",
             not any("VISTA DE PROVA" in (a.get("descricao") or "") for a in res2.atividades), res2.atividades)
    conferir("T226 · as colunas de vista unica ficam NULAS (a vista e sessao)",
             all(a["data_vista_prova"] is None for a in res2.avaliacoes))
    conferir("FR-045 · «(FISCAL)» preenche o fiscal, e o fiscal nao vira o aplicador",
             prova["PM2"]["fiscal"] == "1" and prova["PM2"]["instrutor"] != "1", prova["PM2"])
    conferir("FR-046 · «A / B» guarda o 2º como veio quando nao casa com o cadastro",
             prova["PM1"]["outros_instrutores"] == ["SR. AUXILIAR DE FORA"], prova["PM1"]["outros_instrutores"])
    conferir("FR-031.8 · o curriculo mudo sobre a PM 2: o tempo da planilha vira SUGESTAO para a turma",
             [(g["avaliacao_curriculo_id"], g["ta_aplicacao"]) for g in res2.sugestoes] == [("av-2", 2)], res2.sugestoes)
    confirmado = {**ref2, "ajustes": [{"avaliacao_curriculo_id": "av-2", "origem": "arbitrado", "confirmado": True, "status": "ativo"}]}
    conferir("FR-031.8 · com o arbitrado confirmado, nenhuma sugestao (a confirmada nao e tocada)",
             resolver(leva2, confirmado, dec2, "l.xlsx", set(), True).sugestoes == [])
    sql2 = sincronizar.sql_da_turma(res2, False, None, ensaio=False)
    conferir("T226 · o SQL grava as sessoes em `vistas_de_avaliacao`, por codigo, e as sugestoes no ajuste da turma",
             "insert into public.vistas_de_avaliacao" in sql2 and "'sugestao_planilha'" in sql2 and "-S'" in sql2)
    atuais = {"aulas": [], "atividades": [], "dias_parados": [],
              "avaliacoes": [{**a, "instrutor": a["instrutor"], "status": "ativo", "na_tela": False} for a in res2.avaliacoes]}
    dif2 = sincronizar.diferencas(res2, atuais)["avaliacoes"]
    conferir("duas rodadas: o mesmo plano sobre o que ja esta no banco da ZERO mudanca", not any(dif2.values()), dif2)
    outra = [{**a, "sessoes": a["sessoes"][:1]} if a["chave_ue"] == "PM1" else a for a in atuais["avaliacoes"]]
    conferir("controle: faltando a sessao 2 no banco, a PM 1 muda (o teste discrimina)",
             sincronizar.diferencas(res2, {**atuais, "avaliacoes": outra})["avaliacoes"]["mudadas"] != [])

    emenda_de_15_10()
    correcoes_de_16_10()

    print(f"\n{'TODAS AS PROVAS PASSARAM' if not FALHAS else f'{len(FALHAS)} PROVA(S) FALHARAM'}")
    return 1 if FALHAS else 0


def emenda_de_15_10() -> None:
    """Emenda de 15/10/2026 a spec 014 (Bernardo Villas Boas): RN-AV-10 a RN-AV-14; decisoes de 16/10/2026 (P10, PE,
    sessoes em `sessoes_de_aplicacao`, instrumento gravado). Casos sinteticos, sem banco."""
    from datetime import date as _date

    from scripts.carga_dsa import regras_de_prova as rp
    from scripts.carga_dsa.resolver import Resolvido, _agrupar_sessoes, _ligar_itens, _vista_de_prova_sem_item

    print("EMENDA DE 15/10/2026 — RN-AV-10 a RN-AV-14")
    conferir("RN-AV-10 · «PM1» = «PM 1» = «PM-1»",
             rp.ler_cod("PM1") == rp.ler_cod("PM 1") == rp.ler_cod("PM-1") == ("PM", 1) and rp.ler_cod("PP") == ("PP", None)
             and rp.cod_colado("pm - 2") == "PM2", (rp.ler_cod("PM 1"), rp.ler_cod("PM-1")))

    def pm(n: int, ta: int = 2) -> dict:
        return {"id": f"pm-{n}", "disciplina_id": "d", "numero": n, "instrumento": "PM", "instrumento_nome": "Prova Mista", "ta_aplicacao": ta}

    def pp(n: int) -> dict:
        return {"id": f"pp-{n}", "disciplina_id": "d", "numero": n, "instrumento": "PP", "instrumento_nome": "Prova Prática", "ta_aplicacao": 2}

    tg = {"id": "tg-1", "disciplina_id": "d", "numero": 1, "instrumento": "TG", "instrumento_nome": "Trabalho em Grupo", "ta_aplicacao": 2}
    ti = {"id": "ti-1", "disciplina_id": "d", "numero": 1, "instrumento": "TI", "instrumento_nome": "Trabalho Individual", "ta_aplicacao": 2}

    def prova(cod: str, data: str, desc: str = "", ta: int = 2) -> dict:
        # `sessoes` NOVA a cada prova: lista compartilhada pelo dicionario-base fazia a vista de um caso aparecer no outro.
        return {"chave_ue": cod, "data_avaliacao": data, "descricao": desc, "ta_inicial": 1, "tempos_consumidos": ta,
                "disciplina": "I", "disciplina_id": "d", "sessoes": []}

    def lig(provas: list[dict], itens: list[dict]) -> dict:
        saida = {}
        for n, v in rp.ligar_ao_curriculo(provas, itens).items():
            saida[n] = v["item"]["id"] if "item" in v else ("a_mais:" + v["a_mais"] if "a_mais" in v else "nao_casa")
        return saida

    # O caso que discrimina (DoD 8): ate 15/10 o numero da DESCRICAO vencia; agora decide o COD.
    conferir("RN-AV-10 · o COD decide e a descricao so desempata: «PM1» com descricao «PM2 – …» liga ao item 1",
             lig([prova("PM1", "2026-04-01", "PM2 – ESTATISTICA")], [pm(1), pm(2)]) == {0: "pm-1"})
    conferir("RN-AV-10 · «PM-1» e «PM 1» ligam ao item 1",
             lig([prova("PM-1", "2026-04-01")], [pm(1), pm(2)]) == {0: "pm-1"} and lig([prova("PM 1", "2026-04-01")], [pm(1)]) == {0: "pm-1"})
    conferir("RN-AV-11 · «PM» sem numero e o UNICO item do instrumento", lig([prova("PM", "2026-04-01")], [pm(1)]) == {0: "pm-1"})
    conferir("RN-AV-11 · varios: pela ordem cronologica de aplicacao (1ª aplicada = nº 1), na ordem da lista ou nao",
             lig([prova("PM", "2026-05-01"), prova("PM", "2026-04-01")], [pm(1), pm(2)]) == {0: "pm-2", 1: "pm-1"})
    conferir("RN-AV-11 · numerada e sem numero misturadas: a sem numero pega o MENOR numero livre (interpretacao)",
             lig([prova("PM2", "2026-04-01"), prova("PM", "2026-05-01")], [pm(1), pm(2)]) == {0: "pm-2", 1: "pm-1"})
    conferir("RN-AV-10 · «T» e o trabalho da disciplina: so TG → TG 1",
             lig([prova("T", "2026-04-01", "TRABALHO")], [tg, pm(1)]) == {0: "tg-1"})
    conferir("RN-AV-10 · «T» com TG e TI: a descricao desempata («TRABALHO INDIVIDUAL» → TI); sem pista, nao casa",
             lig([prova("T", "2026-04-01", "TRABALHO INDIVIDUAL - TEMA")], [tg, ti]) == {0: "ti-1"}
             and lig([prova("T", "2026-04-01", "TRABALHO")], [tg, ti]) == {0: "nao_casa"})
    conferir("RN-AV-10 · o numero do «T» e a ORDEM dos trabalhos, nao o item: T1 «INDIVIDUAL» → TI 1; T2 e T3 «EM GRUPO» → TG 1 e TG 2 (CAHO VI)",
             lig([prova("T1", "2026-04-23", "TRABALHO INDIVIDUAL"), prova("T2", "2026-06-23", "TRABALHO EM GRUPO"),
                  prova("T3", "2026-07-06", "TRABALHO EM GRUPO, UE 6")],
                 [ti, tg, {**tg, "id": "tg-2", "numero": 2}]) == {0: "ti-1", 1: "tg-1", 2: "tg-2"})
    conferir("RN-AV-10 · «AV»: vale o instrumento que a DESCRICAO nomeia («PM2 –» → PM 2; «PROVA PRATICA» → PP)",
             lig([prova("AV1", "2026-04-01", "PM2 – FISICA")], [pm(1), pm(2)]) == {0: "pm-2"}
             and lig([prova("AV2", "2026-04-01", "PROVA PRÁTICA DE NAVEGACAO")], [pm(1), pp(1)]) == {0: "pp-1"})
    conferir("D6.1 · «AV» sem descricao que nomeie NAO casa por eliminacao, nem numa disciplina so de PM",
             lig([prova("AV1", "2026-04-01", "AVALIACAO DA APRENDIZAGEM")], [pm(1)]) == {0: "nao_casa"})
    conferir("RN-AV-10 · COD de instrumento que a disciplina nao tem (PP numa disciplina so de PM) nao casa",
             lig([prova("PP", "2026-04-01")], [pm(1)]) == {0: "nao_casa"})
    conferir("RN-AV-14 · numero alem do curriculo, ou o mesmo item pedido duas vezes: «a mais», nao «nao casa»",
             lig([prova("PP2", "2026-04-01")], [pp(1)]) == {0: "a_mais:PP"}
             and lig([prova("PP1", "2026-04-01"), prova("PP1", "2026-06-01")], [pp(1)]) == {0: "pp-1", 1: "a_mais:PP"})

    conferir("RN-AV-11 · «PM» sem numero e descricao «PM4 – …»: vale a ORDEM cronologica, nao o 4 da descricao (so desempata)",
             lig([prova("PM", "2026-03-01", "PM4 – FISICA"), prova("PM", "2026-03-20", "PM1 – X")], [pm(1), pm(2), pm(3), pm(4)]) == {0: "pm-1", 1: "pm-2"})
    conferir("D6.1 · com «AV» o numero ESCRITO na descricao continua valendo («AV1» com «PM4 – …» → PM 4)",
             lig([prova("AV1", "2026-03-01", "PM4 – FISICA")], [pm(1), pm(4)]) == {0: "pm-4"})

    grupos = rp.agrupar_sessoes([prova("PP", d, "PROVA PRATICA DE CAMPO") for d in ("2026-06-22", "2026-06-23", "2026-06-26", "2026-08-19")]
                                + [prova("PP", "2026-06-24", "OUTRA PROVA")], 7)
    conferir("RN-AV-12 · mesma disciplina, mesmo COD e mesma descricao em dias proximos: UMA prova (3 sessoes); 54 dias depois, outra; outra descricao, outra",
             sorted(len(g) for g in grupos) == [1, 1, 3], [[p["data_avaliacao"] for p in g] for g in grupos])
    conferir("RN-AV-12 · a janela e parametro: com 1 dia, o intervalo de 3 dias separa",
             sorted(len(g) for g in rp.agrupar_sessoes([prova("PP", d, "X") for d in ("2026-06-22", "2026-06-23", "2026-06-26")], 1)) == [1, 2])

    r = Resolvido(turma={}, autor={}, titulo="x")
    base = {"tipo_avaliacao": "Prova Prática", "avaliacao_curriculo_id": None, "item": None}
    r.avaliacoes = [{**base, **prova("PP", d, "PROVA PRATICA DE CAMPO", 3), "codigo": f"V{n}"}
                    for n, d in enumerate(("2026-06-22", "2026-06-23", "2026-06-25"))]
    _agrupar_sessoes(r, 7)
    _ligar_itens(r, {"itens_curriculo": [pp(1)]})
    conferir("RN-AV-12 · a PP em 3 dias liga UMA vez ao item; as sessoes 2 e 3 nao sao «a mais» nem erro",
             [a.get("avaliacao_curriculo_id") for a in r.avaliacoes] == ["pp-1", None, None]
             and not any("a mais" in e or "nao casa" in e for e in r.erros_de_planilha)
             and sum("sessao" in i for i in r.informativos) == 2, (r.avaliacoes, r.erros_de_planilha, r.informativos))
    conferir("RN-AV-12 · o TA da prova de 3 sessoes e a SOMA (9) contra o curriculo (2): vai para a lista como TA diferente",
             any("TA de prova diferente — planilha 9, curriculo 2" in e for e in r.erros_de_planilha), r.erros_de_planilha)
    # Desenho aprovado em 16/10/2026 (opcao A): as sessoes 2..N saem de `avaliacoes` e vao para `sessoes_de_aplicacao`.
    from scripts.carga_dsa.resolver import _sessoes_viram_aplicacoes
    for n, a in enumerate(r.avaliacoes):
        a.update({"codigo": f"DSAP-X-L{100 + n:03d}-V", "local": "CAMPO", "instrutor_responsavel_id": None, "aplicacoes": []})
    _sessoes_viram_aplicacoes(r)
    conferir("RN-AV-12 · DISCRIMINA (16/10/2026): a PP em 3 dias e UMA avaliacao; as sessoes 2 e 3 viram `aplicacoes` (…-SA), nao linhas",
             len(r.avaliacoes) == 1
             and [(x["codigo"], x["numero"], x["data"]) for x in r.avaliacoes[0]["aplicacoes"]]
             == [("DSAP-X-L101-SA", 2, "2026-06-23"), ("DSAP-X-L102-SA", 3, "2026-06-25")], r.avaliacoes)

    a_mais = Resolvido(turma={}, autor={}, titulo="x")
    a_mais.avaliacoes = [{**base, **prova("PP1", "2026-04-01", "PRATICA 1"), "codigo": "A"},
                         {**base, **prova("PP2", "2026-09-01", "PRATICA 2"), "codigo": "B"}]
    _agrupar_sessoes(a_mais, 7)
    _ligar_itens(a_mais, {"itens_curriculo": [pp(1)]})
    conferir("RN-AV-14 · prova a mais NAO e erro de planilha: vira informacao «curriculo preve 1 PP; aplicadas 2»",
             not any("a mais" in e for e in a_mais.erros_de_planilha)
             and any("curriculo preve 1 PP; aplicadas 2" in i for i in a_mais.informativos), (a_mais.erros_de_planilha, a_mais.informativos))

    # Revisao de 15/10/2026: a «a mais» que NAO passa do previsto (o COD pede um numero que o curriculo nao tem) nao pode
    # sumir. Caso real: CAHO XIX «PP2» com o curriculo so com PP 3 — 1 aplicada, 1 prevista.
    numero_errado = Resolvido(turma={}, autor={}, titulo="x")
    numero_errado.avaliacoes = [{**base, **prova("PP2", "2026-09-17", "TFM - PROVA PRATICA"), "codigo": "X"}]
    _agrupar_sessoes(numero_errado, 7)
    _ligar_itens(numero_errado, {"itens_curriculo": [{**pp(3), "numero": 3}]})
    conferir("RN-AV-10 · «PP2» com o curriculo so com PP 3 (1 aplicada, 1 prevista): NAO some — vai para a lista e para as provas sem casar",
             any("PP2" in e and "nao casa" in e and "PP 3" in e for e in numero_errado.erros_de_planilha)
             and len(numero_errado.provas_sem_casar) == 1 and not numero_errado.informativos,
             (numero_errado.erros_de_planilha, numero_errado.provas_sem_casar, numero_errado.informativos))
    repetida = Resolvido(turma={}, autor={}, titulo="x")
    repetida.avaliacoes = [{**base, **prova("PM1", "2026-03-01", "A"), "codigo": "A"}, {**base, **prova("PM1", "2026-05-01", "B"), "codigo": "B"}]
    _agrupar_sessoes(repetida, 7)
    _ligar_itens(repetida, {"itens_curriculo": [pm(1), pm(2), pm(3)]})
    conferir("RN-AV-10 · a mesma «PM1» duas vezes, longe, com 3 PM no curriculo: a 2ª «nao casa» (o item PM 1 ja foi ligado), nao some",
             any("ja foi ligado" in e for e in repetida.erros_de_planilha) and len(repetida.provas_sem_casar) == 1, repetida.erros_de_planilha)

    from scripts.carga_dsa.resolver import _casar_vistas
    vg = Resolvido(turma={}, autor={}, titulo="x")
    vg.avaliacoes = [{**base, **prova("PP", d, "PRATICA"), "codigo": f"G{n}", "instrutor": None, "instrutor_responsavel_id": None,
                      "fiscal_id": None, "nome_fiscal_externo": None} for n, d in enumerate(("2026-06-01", "2026-06-05"))]
    _agrupar_sessoes(vg, 7)

    class _Bloco:
        def __init__(self) -> None:
            self.data, self.ta_inicial, self.ue, self.linha, self.tempos, self.chave = _date(2026, 6, 10), 1, "VP", 99, 1, ("I", "VP")
            self.cod = "I"

    class _Linha:
        instrutor = ""

    class _Leitura:
        digitadas: dict = {}
        catalogo = {("I", "VP"): _Linha()}
        datas_avaliacoes: list = []

    pend_vg: list = []
    _casar_vistas(vg, [(_Bloco(), {"cod": "I"})], {"vistas": [{"disciplina": "I", "data": "2026-06-10", "da_prova_de": "2026-06-05"}]},
                  lambda *a, **k: pend_vg.append(a), lambda b, regra=None: None, lambda *a, **k: None, _Leitura())
    conferir("RN-AV-12 · a decisao de vista que aponta a data da 2ª SESSAO vai para a prova (a 1ª sessao), sem pendencia",
             not pend_vg and len(vg.avaliacoes[0]["sessoes"]) == 1 and vg.avaliacoes[1]["sessoes"] == [], (pend_vg, vg.avaliacoes))

    from scripts.carga_dsa.resolver import _sessoes_herdam_o_aplicador
    ha = Resolvido(turma={}, autor={}, titulo="x")
    papel = {"instrutor": None, "instrutor_responsavel_id": None, "fiscal_id": None, "nome_fiscal_externo": None}
    ha.avaliacoes = [{**base, **prova("PP", "2026-06-01", "P"), **papel, "codigo": "H1", "instrutor": "90", "instrutor_responsavel_id": "i-90"},
                     {**base, **prova("PP", "2026-06-02", "P"), **papel, "codigo": "H2", "_aplicador_escrito": False},
                     {**base, **prova("PP", "2026-06-03", "P"), **papel, "codigo": "H3", "_aplicador_escrito": True}]
    _agrupar_sessoes(ha, 7)
    _sessoes_herdam_o_aplicador(ha)
    conferir("RN-AV-12 · a sessao SEM nome na linha herda o aplicador da 1ª; a que traz um nome fora do cadastro fica sem (nao se inventa autoria)",
             [a["instrutor_responsavel_id"] for a in ha.avaliacoes] == ["i-90", "i-90", None], ha.avaliacoes)

    vista = Resolvido(turma={}, autor={}, titulo="x")
    outra = {"disciplina": "IX", "disciplina_id": "d-9"}
    vista.avaliacoes = [{**base, **prova("PP", "2026-04-01", "PROVA PRATICA"), "codigo": "P", **outra},
                        {**base, **prova("T", "2026-04-01", "TRABALHO"), "codigo": "T", **outra},
                        {**base, **prova("AV1", "2026-04-01", "AVALIACAO"), "codigo": "A", **outra},
                        {**base, **prova("PM1", "2026-04-01", "PROVA MISTA"), "codigo": "M", **outra}]
    _vista_de_prova_sem_item(vista, _date(2026, 4, 20), COM_VISTA)
    listadas = sorted(e.split(" ")[1] for e in vista.erros_de_planilha if "vista nao lancada" in e)
    # O caso que discrimina: ate 15/10 as quatro iam para a lista; agora PP e T (sem vista obrigatoria) nao vao.
    conferir("RN-AV-13 · sem item e sem vista no 8º dia: PP e T NAO vao para a lista; PM vai; AV (instrumento desconhecido) vai",
             listadas == ["AV1", "PM1"], listadas)
    pe = Resolvido(turma={}, autor={}, titulo="x")
    pe.avaliacoes = [{**base, **prova("PE1", "2026-04-01", "PROVA ESCRITA"), "codigo": "E", **outra}]
    _vista_de_prova_sem_item(pe, _date(2026, 4, 20), COM_VISTA)
    pe_antes = Resolvido(turma={}, autor={}, titulo="x")
    pe_antes.avaliacoes = [{**base, **prova("PE1", "2026-04-01", "PROVA ESCRITA"), "codigo": "E", **outra}]
    _vista_de_prova_sem_item(pe_antes, _date(2026, 4, 20), frozenset({"PM", "PO"}))
    conferir("RN-AV-13 · DISCRIMINA (16/10/2026, «PM, PO e PE»): a PE sem item e sem vista vai para a lista; com a marca antiga, nao",
             any("vista nao lancada" in e for e in pe.erros_de_planilha) and not pe_antes.erros_de_planilha,
             (pe.erros_de_planilha, pe_antes.erros_de_planilha))
    conferir("RN-AV-13 · a marca e DADO do banco: `vista_obrigatoria` decide pelo conjunto que chega (o JSON nao tem mais a lista)",
             [rp.vista_obrigatoria(i, COM_VISTA) for i in ("PM", "PO", "PE", "PP", "TG", "TI")] == [True, True, True, False, False, False]
             and "vistaObrigatoria" not in rp.REGRAS and not hasattr(rp, "VISTA_OBRIGATORIA"))

    # P10 (aceita em 16/10/2026): com «AV» e a descricao muda, a T/E da linha desempata.
    av3 = {**prova("AV3", "2026-09-01", "AV3 – AVALIACAO - PERMANENCIA NA AGUA")}
    conferir("P10 · DISCRIMINA: «AV3» sem instrumento na descricao, T/E «PP», curriculo so com PP 3 → liga a PP 3; sem a T/E, nao casa",
             lig([{**av3, "_tecnica": "PP"}], [{**pp(3), "id": "pp-3", "numero": 3}]) == {0: "pp-3"}
             and lig([av3], [{**pp(3), "id": "pp-3", "numero": 3}]) == {0: "nao_casa"})
    conferir("P10 · a descricao que NOMEIA vence a T/E («PM2 – …» com T/E «PP» → PM 2)",
             lig([{**prova("AV1", "2026-04-01", "PM2 – FISICA"), "_tecnica": "PP"}], [pm(1), pm(2), pp(1)]) == {0: "pm-2"})
    conferir("P10 · T/E que nao e instrumento da disciplina (EO) nao desempata",
             lig([{**prova("AV1", "2026-04-01", "AVALIACAO"), "_tecnica": "EO"}], [pm(1)]) == {0: "nao_casa"})

    # A2 (desenho aprovado em 16/10/2026): o instrumento gravado da prova SEM item.
    nomes = {"PE": "Prova Escrita", "PM": "Prova Mista", "PP": "Prova Prática", "PO": "Prova Oral", "TG": "Trabalho em Grupo",
             "TI": "Trabalho Individual", "OD": "Observação de Desempenho"}
    conferir("A2 · o instrumento da prova sem item: o COD que nomeia um so («PP» → PP); «T» desempatado pela descricao; ambiguo → nenhum",
             rp.instrumento_da_prova(prova("PP", "2026-06-22", "PROVA PRATICA DE CAMPO"), nomes) == "PP"
             and rp.instrumento_da_prova(prova("T", "2026-07-14", "TRABALHO INDIVIDUAL"), nomes) == "TI"
             and rp.instrumento_da_prova(prova("T", "2026-07-14", "TRABALHO"), nomes) is None)
    conferir("A2 · «AV»: pela descricao, depois pela T/E (P10); nada nomeado → nenhum (D6.1: nunca por eliminacao)",
             rp.instrumento_da_prova(prova("AV1", "2026-04-01", "PROVA ORAL DE INGLES"), nomes) == "PO"
             and rp.instrumento_da_prova({**prova("AV2", "2026-04-01", "AVALIACAO"), "_tecnica": "PM"}, nomes) == "PM"
             and rp.instrumento_da_prova(prova("AV2", "2026-04-01", "AVALIACAO"), nomes) is None)


def correcoes_de_16_10() -> None:
    """As correcoes da lista de erros decididas por Bernardo Villas Boas em 16/10/2026 (noite): R1, R2, R4, R5, R7, a
    sugestao vazia (PEND-014-11) e as decisoes por LINHA (item ligado, vista ligada e vista dividida). Sem banco.
    Cada caso diz o veredito de ANTES ao lado, que e o que ele discrimina (DoD 8)."""
    import tempfile
    from datetime import date as _date

    from scripts.carga_dsa.resolver import Resolvido, _casar_vistas, _ligar_itens, _sugestoes

    print("CORRECOES DE 16/10/2026 — a lista de erros (R1, R2, R4, R5, R7, sugestao vazia, decisoes por linha)")

    def pm(n: int, ta: int = 2) -> dict:
        return {"id": f"pm-{n}", "disciplina_id": "d", "numero": n, "instrumento": "PM", "instrumento_nome": "Prova Mista",
                "ta_aplicacao": ta, "origem_ta_aplicacao": "curriculo", "origem_ta_vista": "curriculo"}

    def pp(n: int, ta: int = 2, vista: bool = True) -> dict:
        return {"id": f"pp-{n}", "disciplina_id": "d", "numero": n, "instrumento": "PP", "instrumento_nome": "Prova Prática",
                "ta_aplicacao": ta, "origem_ta_aplicacao": "curriculo", "origem_ta_vista": "curriculo" if vista else None}

    def prova(cod: str, data: str, codigo: str, linha: int, desc: str = "AVALIACAO", ta: int = 2, inicio: int = 1) -> dict:
        return {"chave_ue": cod, "cod_planilha": "I", "data_avaliacao": data, "descricao": desc, "ta_inicial": inicio,
                "tempos_consumidos": ta, "disciplina": "I", "disciplina_id": "d", "sessoes": [], "codigo": codigo, "linha": linha,
                "observacoes": "procedencia.", "instrutor": None, "instrutor_responsavel_id": None, "fiscal_id": None,
                "nome_fiscal_externo": None, "aplicacoes": []}

    # ── decisao por linha: a prova liga ao item que o lote nomeia (as 17, CAHO-01, C-Espc-HN-01) ──
    def ligar(provas: list[dict], itens: list[dict], decisoes: dict | None = None, atuais: list[dict] | None = None) -> Resolvido:
        r = Resolvido(turma={}, autor={}, titulo="x")
        r.avaliacoes = provas
        _ligar_itens(r, {"itens_curriculo": itens, "atuais": {"avaliacoes": atuais or []}}, decisoes)
        return r

    sem = ligar([prova("AV1", "2026-04-01", "X1", 10)], [pm(1), pm(2)])
    com = ligar([prova("AV1", "2026-04-01", "X1", 10)], [pm(1), pm(2)], {"itens_por_prova": {"10": {"item": "PM 2", "motivo": "TESTE-01"}}})
    conferir("16/10 · DISCRIMINA: a decisao `itens_por_prova` liga a prova da LINHA 10 a PM 2 (sem ela, «AV1» sem instrumento nao casa)",
             sem.avaliacoes[0].get("avaliacao_curriculo_id") is None and com.avaliacoes[0]["avaliacao_curriculo_id"] == "pm-2"
             and "por decisao: TESTE-01" in com.avaliacoes[0]["observacoes"], (sem.avaliacoes, com.avaliacoes))
    duas = ligar([prova("PM", "2026-04-01", "X1", 10), prova("PM", "2026-05-01", "X2", 20)], [pm(1), pm(2)],
                 {"itens_por_prova": {"10": {"item": "PM 2", "motivo": "T"}}})
    conferir("16/10 · o item decidido sai da ligacao automatica: a outra «PM» fica com a PM 1, e nada vira «a mais»",
             [a["avaliacao_curriculo_id"] for a in duas.avaliacoes] == ["pm-2", "pm-1"] and not duas.informativos, duas.avaliacoes)

    # ── R7: a lista olha o que o sistema ja gravou ──
    pp36 = ligar([prova("PP", "2026-06-22", "X3", 30, ta=36)], [pp(1, ta=40)])
    pp36_juntada = ligar([prova("PP", "2026-06-22", "X3", 30, ta=36)], [pp(1, ta=40)],
                         atuais=[{"codigo": "X3", "ta_fora_da_planilha": 4}])
    conferir("R7 · DISCRIMINA: «planilha 36 × curriculo 40» some quando o sistema tem +4 TA de sessao nascida na tela (a juncao da CAHO IV)",
             any("TA de prova diferente" in e for e in pp36.erros_de_planilha)
             and not any("TA de prova diferente" in e for e in pp36_juntada.erros_de_planilha)
             and any("R7" in i for i in pp36_juntada.informativos), (pp36.erros_de_planilha, pp36_juntada.erros_de_planilha))
    duas_pp = [prova("PP", "2026-06-22", "X3", 30), prova("PP", "2026-08-19", "X4", 40)]
    antes = ligar([dict(x) for x in duas_pp], [pp(1)])
    depois = ligar([dict(x) for x in duas_pp], [pp(1)], atuais=[{"codigo": "X4", "juntada": True}])
    conferir("R7 · DISCRIMINA: a prova JUNTADA como sessao nao conta de novo — «curriculo preve 1 PP; aplicadas 2» some",
             any("aplicadas 2" in i for i in antes.informativos) and not any("aplicadas 2" in i for i in depois.informativos),
             (antes.informativos, depois.informativos))
    na_tela = ligar([prova("AV1", "2026-04-01", "X5", 50)], [pm(1)], atuais=[{"codigo": "X5", "na_tela": True, "item": "PM 1"}])
    conferir("R7 · a prova que nao casa na planilha, mas esta LIGADA na tela, nao vai para a lista",
             not na_tela.erros_de_planilha and not na_tela.provas_sem_casar, na_tela.erros_de_planilha)

    # ── PEND-014-11: sugestao vazia ──
    sv = Resolvido(turma={}, autor={}, titulo="x")
    sv.avaliacoes = [{**prova("PP", "2026-06-22", "X6", 60), "avaliacao_curriculo_id": "pp-1", "item": "PP 1"}]
    _sugestoes(sv, {"itens_curriculo": [pp(1, vista=False)]})
    sv_com_vista = Resolvido(turma={}, autor={}, titulo="x")
    sv_com_vista.avaliacoes = [{**prova("PP", "2026-06-22", "X6", 60), "avaliacao_curriculo_id": "pp-1", "item": "PP 1",
                                "sessoes": [{"tempos_consumidos": 1}]}]
    _sugestoes(sv_com_vista, {"itens_curriculo": [pp(1, vista=False)]})
    conferir("PEND-014-11 · DISCRIMINA: curriculo mudo sobre a vista e planilha SEM vista → nenhuma sugestao (antes: uma, com os dois tempos nulos)",
             sv.sugestoes == [] and [g["ta_vista"] for g in sv_com_vista.sugestoes] == [1], (sv.sugestoes, sv_com_vista.sugestoes))

    # ── R2 e as decisoes de vista por linha ──
    class _B:
        def __init__(self, cod: str, ue: str, dia: _date, ta: int, linha: int, tempos: int = 1) -> None:
            self.cod, self.ue, self.data, self.ta_inicial, self.linha, self.tempos = cod, ue, dia, ta, linha, tempos
            self.chave = (cod, ue)

    class _L:
        instrutor = ""

    class _Leitura:
        def __init__(self, datas: list[dict]) -> None:
            self.digitadas: dict = {}
            self.catalogo = _Qualquer()
            self.datas_avaliacoes = datas

    class _Qualquer(dict):
        def __getitem__(self, chave: object) -> _L:
            return _L()

    def casar(provas: list[dict], blocos: list[_B], decisoes: dict, datas: list[dict] | None = None) -> tuple[Resolvido, list]:
        r = Resolvido(turma={}, autor={}, titulo="x")
        r.avaliacoes = provas
        pend: list = []
        _casar_vistas(r, [(b, {"cod": "I"}) for b in blocos], decisoes, lambda *a, **k: pend.append(a), lambda b, regra=None: None,
                      lambda *a, **k: None, _Leitura(datas or []), codigo=lambda b, suf: f"DSAP-T-L{b.linha:03d}-{suf}")
        return r, pend

    ti, po = prova("TI", "2026-03-12", "TI", 511, inicio=1, ta=1), prova("PO", "2026-03-12", "PO", 512, inicio=2, ta=1)
    datas = [{"cod": "I", "chave_prova": "ITI", "chave_vista": "IVP2", "data_prova": _date(2026, 3, 12), "data_vista": _date(2026, 3, 12)},
             {"cod": "I", "chave_prova": "IPO", "chave_vista": "IVP3", "data_prova": _date(2026, 3, 12), "data_vista": _date(2026, 3, 13)}]
    r2, _p = casar([ti, po], [_B("I", "VP2", _date(2026, 3, 12), 7, 517), _B("I", "VP3", _date(2026, 3, 13), 7, 526)], {}, datas)
    conferir("R2 · DISCRIMINA: duas provas no mesmo dia — a aba de datas manda a VP2 para a TI e a VP3 para a PO (antes: as duas para a TI)",
             [s["codigo"] for s in ti["sessoes"]] == ["DSAP-T-L517-S"] and [s["codigo"] for s in po["sessoes"]] == ["DSAP-T-L526-S"],
             (ti["sessoes"], po["sessoes"]))
    av2, av1 = prova("AV2", "2026-03-12", "B2", 37, inicio=7, ta=3), prova("AV1", "2026-03-13", "B1", 42, inicio=3, ta=3)
    casar([av2, av1], [_B("I", "VP1", _date(2026, 3, 13), 1, 40), _B("I", "VP", _date(2026, 3, 13), 6, 45)], {})
    conferir("R2 · DISCRIMINA (BATI): a VP1 do 1º tempo NAO vai para a AV1 do 3º tempo do mesmo dia — vai para a AV2 de 12/03; a VP, para a AV1",
             [s["ta_inicial"] for s in av2["sessoes"]] == [1] and [s["ta_inicial"] for s in av1["sessoes"]] == [6],
             (av2["sessoes"], av1["sessoes"]))
    t2, pm3 = prova("T2", "2026-06-23", "T2", 1452), prova("PM3", "2026-07-02", "M3", 1530)
    _r, pend_vi = casar([t2, pm3], [_B("I", "VP3", _date(2026, 6, 24), 6, 1466), _B("I", "VT2", _date(2026, 7, 1), 9, 1529)], {})
    conferir("R2 · DISCRIMINA (CAHO VI): a VP3 anterior a PM3 NAO cai no T2 que a VT2 reivindica — fica sem prova; a VT2 vai para o T2",
             [s["codigo"] for s in t2["sessoes"]] == ["DSAP-T-L1529-S"] and pm3["sessoes"] == []
             and [x[0] for x in pend_vi] == ["vista_sem_prova"], (t2["sessoes"], pend_vi))
    a1042, b1148 = prova("AV1", "2026-05-13", "P1", 1042), prova("AV2", "2026-05-25", "P2", 1148, inicio=5, ta=3)
    casar([a1042, b1148], [_B("I", "VP", _date(2026, 5, 27), 1, 1162, tempos=2)],
          {"vistas": [{"linha": 1162, "dividir": [{"tempos": 1, "da_prova_linha": 1042}, {"tempos": 1, "da_prova_linha": 1148}]}]})
    conferir("16/10 (V2) · a decisao DIVIDE o bloco de vista de 2 TA: 1º tempo para a prova de 13/05, 2º para a de 25/05, cada um com o codigo da sua linha",
             [(s["codigo"], s["ta_inicial"], s["tempos_consumidos"]) for s in a1042["sessoes"]] == [("DSAP-T-L1162-S", 1, 1)]
             and [(s["codigo"], s["ta_inicial"], s["tempos_consumidos"]) for s in b1148["sessoes"]] == [("DSAP-T-L1163-S", 2, 1)],
             (a1042["sessoes"], b1148["sessoes"]))
    t1636, t1743 = prova("T", "2026-07-14", "T1", 1636, inicio=5, ta=1), prova("T", "2026-07-27", "T2", 1743, inicio=1, ta=3)
    casar([t1636, t1743], [_B("I", "VT", _date(2026, 7, 27), 4, 1746)], {"vistas": [{"linha": 1746, "da_prova_linha": 1636}]})
    conferir("16/10 (CAHO-06) · DISCRIMINA: a decisao por LINHA manda os comentarios de 27/07 para a TI de 14/07 (sem ela, iam para a de 27/07)",
             len(t1636["sessoes"]) == 1 and t1743["sessoes"] == [], (t1636["sessoes"], t1743["sessoes"]))

    # ── R4 e R1: pela planilha ──
    with tempfile.TemporaryDirectory() as pasta:
        livro = openpyxl.Workbook()
        cat = livro.active
        cat.title = planilha.ABA_CATALOGO
        cat.append([None])
        cat.append(["CÓD", "DISCIPLINA", "Nº UE", "CONCATENADO", "UNIDADES DE ENSINO E TÓPICOS", "CH", "LOCAL", "T/E", "INSTRUTOR/PROFESSOR"])
        for linha in (["I", "DISCIPLINA UM", "1P", "I1P", "1 – TOPICO DA PLANILHA", 1.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
                      ["I", "DISCIPLINA UM", 1.0, "I1", "2 – SEGUNDO TOPICO", 1.0, "SALA 2", "EO", "CC (T) FULANO SILVA"],
                      ["IV", "VISITA", 2.0, "IVVT", "VISITA AO FAROL", 8.0, "MACAE", "AP", "--"],
                      ["CÓD", "ORDEM", "CONCATENADO", None, None, None, None, None, None]):
            cat.append(linha)
        folha = livro.create_sheet(planilha.ABA_PREENCHIMENTO)
        folha["A1"], folha["B1"] = "SIGLA DO CURSO:", "C-Exp-QUALQUER"
        folha["A2"], folha["B2"] = "Nº DE ALUNOS:", 6
        tas = {1: ("I", "1P"), 2: ("I", 1.0), 4: ("IV", "VT")}
        for n in range(1, 10):
            folha.cell(3 + n, 3, f"{n}º")
            if n == 3:
                folha.cell(3 + n, 2, datetime(2026, 3, 11))
            if n in tas:
                folha.cell(3 + n, 4, tas[n][0])
                folha.cell(3 + n, 5, tas[n][1])
        livro.create_sheet(planilha.ABA_CONTROLE).append(["CÓD.", "DISCIPLINA", "CH. PREVISTA", "CH. CUMPRIDA"])
        caminho = Path(pasta) / "r1r4.xlsx"
        livro.save(caminho)
        leitura = planilha.ler(caminho)
    conferir("R4 · DISCRIMINA: a chave e a coluna D, como o PROCV — «IV»+«VT» acha a linha «IVVT» (por A + C, «IV»+«2», ela ficava fora)",
             ("IV", "VT") in leitura.catalogo and ("IV", "2") not in leitura.catalogo and ("IV", "VT") not in leitura.sem_catalogo,
             (sorted(leitura.catalogo), leitura.sem_catalogo))
    conferir("R4 · a lista aponta A + C ≠ D (o cabecalho repetido no meio da aba nao conta)",
             getattr(leitura.catalogo, "divergentes", None) == [(5, "IV2", "IVVT")], getattr(leitura.catalogo, "divergentes", None))
    ref = retrato()
    ref["ues"].append({"id": "u-2", "codigo": "UE-2", "disciplina_id": "d-1", "numero": 2, "topico": "SEGUNDO TOPICO", "ch": 1})
    dec = {**DECISOES, "chaves": {**DECISOES["chaves"], "IV+VT": {"tipo": "ignorar"}}}
    sem_de_para = resolver(leitura, ref, dec, "r1.xlsx", set(), True)
    com_de_para = resolver(leitura, ref, {**dec, "ues": {"I+1P": "I+1", "I+1": "I+2"}}, "r1.xlsx", set(), True)
    conferir("R1 · DISCRIMINA: com o de-para, «I+1P» fica na UE 1 e «I+1» vai para a UE 2 (sem ele, as duas na UE 1)",
             [a["ue"] for a in sem_de_para.aulas] == ["UE-1", "UE-1"] and [a["ue"] for a in com_de_para.aulas] == ["UE-1", "UE-2"],
             ([a["ue"] for a in sem_de_para.aulas], [a["ue"] for a in com_de_para.aulas]))
    conferir("R1 · a chave CRUA vem primeiro: so com «I+1» → «I+2» no de-para, a «I+1P» iria (errado) para a UE 2 — o de-para tem de trazer a «1P»",
             [a["ue"] for a in resolver(leitura, ref, {**dec, "ues": {"I+1": "I+2"}}, "r1.xlsx", set(), True).aulas] == ["UE-2", "UE-2"])
    ch_ue = [e for e in com_de_para.erros_tabelados if e["diverge"].startswith("CH da UE")]
    conferir("R1 · a CH por UE compara pela UE do de-para, e duas linhas do catalogo na MESMA UE somam (1 + 1 = 2 contra 3 na UE 1)",
             [(e["diverge"], e["planilha"]) for e in resolver(leitura, ref, {**dec, "ues": {"I+1P": "I+1", "I+1": "I+1"}}, "r1.xlsx",
                                                               set(), True).erros_tabelados if e["diverge"].startswith("CH da UE")]
             == [("CH da UE 1", "2 TA (BD DISCIPLINAS)")]
             and [(e["diverge"], e["planilha"]) for e in ch_ue] == [("CH da UE 1", "1 TA (BD DISCIPLINAS)")], ch_ue)

    # ── as listas: o que ja foi decidido e as traducoes novas (erros_por_turma) ──
    from scripts.carga_dsa import decididos
    from scripts.carga_dsa import erros_por_turma as ept
    item = decididos.carregar()["itens"][0]
    t = ("1", "Provas", item["onde"], item["oque"], "Conferir.")
    usados: set[str] = set()
    conferir("16/10 · o item que Bernardo JA DECIDIU sai como informacao, com a decisao e o ID; o resto nao muda",
             ept.decidir(item["turma"], t, usados)[0] == ept.INFORMACAO and item["id"] in ept.decidir(item["turma"], t, usados)[4]
             and ept.decidir(item["turma"], ("1", "Provas", item["onde"], item["oque"] + " x", "Conferir."), set())[0] == "1")
    conferir("R8 · o aviso de tempo do TFM (C-Ap-FR X) e informacao; o mesmo aviso em outra disciplina, nao",
             (decididos.do_item("C-Ap-FR 2026", "Disciplina X, AV9 de 01/09/2026", "Tempo da prova diferente do currículo: planilha 3 TA, currículo 2 TA (PP 1).") or {}).get("id") == "R8"
             and decididos.do_item("C-Ap-FR 2026", "Disciplina IX, AV9 de 01/09/2026", "Tempo da prova diferente do currículo: planilha 3 TA, currículo 2 TA (PP 1).") is None)
    r4 = ept.traduzir("T", ["—", "BD DISCIPLINAS linha 59: a chave por A + C «IV2» difere da coluna D «IVVT» — vale a D, como o PROCV da IMPRESSAO", "—", "—", "x"], {})
    r7 = ept.traduzir("T", ["—", "IV PP de 2026-06-22: a planilha nao casa (x), mas no sistema a prova ja esta ligada ou juntada — resolvido no sistema (R7)", "—", "—", "x"], {})
    conferir("R4 · a divergencia A + C × D entra na lista, prioridade baixa; R7 · o que o sistema ja tem e informacao (antes: «nao casa», prioridade alta)",
             r4 is not None and r4[0] == 6 and "IVVT" in r4[3] and r7 is not None and r7[0] == ept.INFORMACAO, (r4, r7))


if __name__ == "__main__":
    raise SystemExit(main())
