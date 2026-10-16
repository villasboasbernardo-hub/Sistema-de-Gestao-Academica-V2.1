"""A lista de erros de planilha da rodada, uma versão por turma, para repassar aos operadores (fora do git).

Uso: `python -X utf8 -m scripts.carga_dsa.erros_por_turma <pasta-da-rodada>`

Lê `erros-de-planilha.md` da rodada (gerado por `scripts.carga_dsa.sincronizar`) e escreve em
`<rodada>/erros-por-turma/` um `.md` por turma e `erros-por-turma.xlsx`, com uma aba por turma e as colunas
**Prioridade** e **Resposta do operador** *(pedido de Bernardo Villas Boas, 15/10/2026)*.

Emenda de 15/10/2026 (spec 014): a **prova a mais** que o currículo (RN-AV-14) e as **sessões seguintes** de uma
mesma prova (RN-AV-12) são informação — não são erro de planilha.
Decisão de Bernardo Villas Boas, 16/10/2026 (P9): *"Catálogo: código inexistente (PROCV) e T/E da avaliação voltam a
"corrigir", prioridade Média; o resto é informação."* — a CH da UE e da disciplina continuam **informação** (o sistema
usa a CH do currículo, D-CUR-1).

Decisões de Bernardo Villas Boas, 16/10/2026 (noite): o item que ele JÁ DECIDIU — «não é erro», «fato», «já
resolvido» e o aviso de tempo do TFM (R8) — sai como **informação**, com «decidido em 16/10/2026 (ID)»
(`decididos.py`, registro `decisoes-da-lista.json`); a divergência A + C × D do catálogo (R4) entra com prioridade baixa.

⚠️ A virada de cada turma vem de `fontes.json` (a `virada` de cada turma), não de uma cópia escrita aqui: a lista não
pede para corrigir na planilha o que a turma já lança no sistema.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from scripts.carga_dsa import decididos

FONTES = Path(__file__).resolve().parent / "fontes.json"
ORDEM = ["CAHO 2026", "C-Ap-FR 2026", "C-Esp-ME 2026", "C-Ap-HN 2026", "C-Espc-FR 2026", "C-Espc-HN 2026",
         "C-Exp-Ag-Mag 2026", "C-Exp-MetocOf 2026", "C-Exp-Obs-ME 2026", "C-Exp-BATI T1 2026"]
INFORMACAO = 9
# 4: o catálogo que é ERRO da planilha (PROCV e T/E da avaliação) — prioridade Média (P9, 16/10/2026).
PRIORIDADE = {1: "1 — alta", 2: "2 — alta", 3: "3 — média", 4: "3 — média", 5: "4 — média", 6: "5 — baixa", 7: "5 — baixa",
              8: "5 — baixa", INFORMACAO: "Informação"}


def viradas() -> dict[str, str]:
    turmas = json.loads(FONTES.read_text(encoding="utf-8")).get("turmas", [])
    return {t["turma"]: t["virada"] for t in turmas if t.get("virada")}


def br(d: str) -> str:
    return f"{d[8:10]}/{d[5:7]}/{d[0:4]}"


def traduzir(turma: str, c: list[str], virada: dict[str, str]) -> tuple | None:
    """(prioridade, assunto, onde, o que esta errado, o que fazer). None = nao entra (depois da virada)."""
    disc, oque, plan, curr = c[0], c[1], c[2], c[3]
    fez = c[4] if len(c) > 4 else ""
    m = re.match(r"(\S+) (.+?) de (\d{4}-\d{2}-\d{2}): (.*)", oque)
    if m:
        d, sig, data, resto = m.groups()
        onde = f"Disciplina {d}, {sig} de {br(data)}"
        if virada.get(turma) and data >= virada[turma]:
            return None  # depois da virada a turma lanca no sistema; a rodada nao le a planilha (prova, vista, TA, RECON, A4.7)
        if "(R7)" in resto:
            # R7 (16/10/2026): o que o sistema ja tem (juncao, ligacao na tela, sessao nascida na tela) — informacao.
            return (INFORMACAO, "Já feito no sistema", onde,
                    resto.replace("curriculo", "currículo").replace("nao casa", "não casa").replace("ja esta", "já está"),
                    "Nada a corrigir.")
        if "sessao" in resto and "RN-AV-12" in resto:
            k = re.search(r"sessao (\d+) da prova de (\d{4}-\d{2}-\d{2})", resto)
            return (INFORMACAO, "Prova em várias sessões (RN-AV-12)", onde,
                    f"É a sessão {k.group(1)} da prova de {br(k.group(2))} (mesma disciplina, mesmo código, mesma descrição)." if k else resto,
                    "Nada a corrigir: é a mesma prova em mais de um dia.")
        if "nao casa" in resto:
            motivo = re.search(r"\((.*)\)\s*$", resto)
            return (1, "Provas", onde,
                    "A prova não corresponde a nenhuma avaliação prevista no currículo da disciplina"
                    + (f" — {motivo.group(1).replace('nao ', 'não ').replace('descricao', 'descrição').replace('tambem', 'também').replace('curriculo', 'currículo')}." if motivo else "."),
                    "Conferir o código da prova na coluna U.E.: ele deve ser o instrumento e o número como no currículo "
                    "(ex.: PM1, PP, TG1). Se for «AV», escrever o instrumento no início da descrição (ex.: «PM2 – …»).")
        if "prova a mais" in resto:
            n = re.search(r"preve (\d+ \S+)", resto)
            return (INFORMACAO, "Prova a mais que o currículo (RN-AV-14)", onde,
                    f"Prova a mais: o currículo prevê {n.group(1) if n else 'menos'}.",
                    "Nada a corrigir: não é erro de planilha (RN-AV-14).")
        if "TA de prova diferente" in resto:
            n = re.search(r"planilha (\d+), curriculo (\d+) \(([^)]+)\)", resto)
            return (1, "Provas", onde,
                    f"Tempo da prova diferente do currículo: planilha {n.group(1)} TA, currículo {n.group(2)} TA ({n.group(3)})." if n else resto,
                    "Conferir o nº de tempos lançados para a prova.")
        if "vista nao lancada" in resto:
            return (2, "Vistas de prova", onde, "A prova não tem vista lançada.",
                    "Lançar a vista de prova (data e tempos) ou informar à Divisão que esta prova não tem vista.")
        if "nao tem avaliacao do curriculo carregada" in resto:
            return (INFORMACAO, "Avaliações do currículo ainda não carregadas", onde,
                    "A disciplina ainda não tem as avaliações do currículo carregadas no sistema.",
                    "Nada a corrigir na planilha: pendente na Divisão (A4.7).")
        if "RECON" in resto:
            return (1, "Provas", onde, "Linha de recuperação (RECON) sem correspondente no currículo.", "Conferir se a recuperação foi aplicada.")
        return (8, "Outros", onde, resto, "Conferir.")
    m = re.match(r"(\S+): curriculo preve (\d+) (\S+); aplicadas (\d+) \(([^)]*)\)", oque)
    if m:
        d, prev, inst, apl, datas = m.groups()
        datas_br = ", ".join(br(x.strip()) for x in datas.split(",") if x.strip())
        return (INFORMACAO, "Prova a mais que o currículo (RN-AV-14)", f"Disciplina {d}",
                f"O currículo prevê {prev} {inst}; foram aplicadas {apl} ({datas_br}).",
                "Nada a corrigir: não é erro de planilha. A Divisão acompanha pela aba Avaliações da turma.")
    m = re.match(r"aula em (\d{4}-\d{2}-\d{2}) nos tempos (\S+), no dia que a mesma planilha marca parado", oque)
    if m:
        return (3, "Aulas em dia parado", f"Disciplina {disc}, {br(m.group(1))}, tempos {m.group(2)}",
                f"Aula lançada num dia que a própria planilha marca como parado ({plan.split(' e parada ')[-1]}).",
                "Conferir: ou a aula não aconteceu (apagar) ou o dia não estava parado (corrigir a parada).")
    m = re.match(r"CH da UE (\d+)", oque)
    if m:
        return (INFORMACAO, "Catálogo (aba BD DISCIPLINAS)", f"Disciplina {disc}, UE {m.group(1)}",
                f"CH da UE na planilha: {plan.replace(' (BD DISCIPLINAS)', '')}; no currículo: {curr}.",
                "Para informação: o sistema usa a CH do currículo.")
    if oque == "CH da disciplina":
        return (INFORMACAO, "Catálogo (aba BD DISCIPLINAS)", f"Disciplina {disc}",
                f"CH da disciplina na planilha: {plan}; no currículo: {curr}.",
                "Para informação: o sistema usa a CH do currículo.")
    m = re.match(r"linha (\d+) \(([^)]*)\): a chave (\S+) nao existe no catalogo", oque)
    if m:
        return (4, "Catálogo (aba BD DISCIPLINAS)", f"Aba PREENCHIMENTO, linha {m.group(1)} ({m.group(2)})",
                f"O código «{m.group(3)}» não existe na aba BD DISCIPLINAS (na IMPRESSÃO sai erro de PROCV).",
                "Incluir o código na aba BD DISCIPLINAS ou corrigir a chave na PREENCHIMENTO: até lá o lançamento fica fora da rodada.")
    m = re.match(r"linha (\d+) \(([^)]*)\): a avaliacao (\S+) esta com T/E «([^»]*)» no catalogo", oque)
    if m:
        return (4, "Catálogo (aba BD DISCIPLINAS)", f"Aba PREENCHIMENTO, linha {m.group(1)} ({m.group(2)})",
                f"A avaliação «{m.group(3)}» está com técnica de ensino «{m.group(4)}» no catálogo.",
                "Escrever na coluna T/E do catálogo o instrumento da avaliação (PM, PP, PE, TG…), e não EO nem vazio.")
    m = re.match(r"BD DISCIPLINAS linha (\d+): a chave por A \+ C «([^»]*)» difere da coluna D «([^»]*)»", oque)
    if m:
        # R4 (16/10/2026): a sincronizacao le a chave pela coluna D, como o PROCV; a lista aponta onde A + C diverge.
        return (6, "Catálogo (aba BD DISCIPLINAS)", f"Aba BD DISCIPLINAS, linha {m.group(1)}",
                f"As colunas A + C formam «{m.group(2)}», e a coluna D (CONCATENADO) diz «{m.group(3)}». "
                "A IMPRESSÃO e o sistema usam a D.",
                "Conferir: se a D estiver certa, acertar o Nº UE (coluna C); se não, corrigir a D.")
    m = re.match(r"UE (\S+): (\d+) TA lancados e (\d+) previstos no catalogo \((\w+)\)", oque)
    if m:
        return (5, "Carga por UE", f"UE {m.group(1)}",
                f"Lançados {m.group(2)} TA; previstos {m.group(3)} TA ({'passou' if m.group(4) == 'PASSOU' else 'falta'}).",
                "Conferir os lançamentos da UE.")
    if "HORARIOS" in oque or "CFG-H" in oque:
        return (6, "Horários", "Aba HORÁRIOS", oque.replace("aba HORARIOS", "aba HORÁRIOS").replace(" as ", " às "),
                "Conferir os horários dos tempos.")
    if oque.startswith("IMPRESSAO"):
        return (7, "Aba IMPRESSÃO", oque.split(":")[0].replace("IMPRESSAO", "Aba IMPRESSÃO"), oque.split(": ", 1)[1],
                "Restaurar a fórmula da célula.")
    if "informacao" in fez:
        return (INFORMACAO, "Informação", disc, oque, "Nada a corrigir.")
    return (8, "Outros", disc, oque, "Conferir.")


def decidir(turma: str, t: tuple, usados: set[str]) -> tuple:
    """O item que Bernardo ja decidiu (16/10/2026) vira informacao, com a decisao e o ID."""
    d = decididos.do_item(turma, t[2], t[3])
    if d is None:
        return t
    usados.add(d["id"])
    return (INFORMACAO, f"Decidido em {decididos.data()}", t[2], t[3], f"Decisão: {d['decisao']} ({d['id']}). Nada a corrigir.")


def ler(rodada: Path, virada: dict[str, str], usados: set[str] | None = None) -> dict[str, list[tuple]]:
    usados = usados if usados is not None else set()
    secoes: dict[str, list[tuple]] = {}
    turma = None
    for linha in (rodada / "erros-de-planilha.md").read_text(encoding="utf-8").splitlines():
        m = re.match(r"## (.+)", linha)
        if m:
            turma = m.group(1).strip()
            secoes[turma] = []
            continue
        if turma and linha.startswith("| ") and not linha.startswith("| disciplina") and not linha.startswith("|---"):
            c = [x.strip() for x in linha.strip().strip("|").split(" | ")]
            if c[1] == "nenhum erro":
                continue
            t = traduzir(turma, c, virada)
            if t:
                secoes[turma].append(decidir(turma, t, usados))
    return secoes


def main() -> int:
    rodada = Path(sys.argv[1])
    saida = rodada / "erros-por-turma"
    saida.mkdir(exist_ok=True)
    virada = viradas()
    usados: set[str] = set()
    secoes = ler(rodada, virada, usados)
    livro = Workbook()
    livro.remove(livro.active)
    print("turma\ta corrigir\tinformação")
    for i, turma in enumerate([t for t in ORDEM if t in secoes] + [t for t in secoes if t not in ORDEM], start=1):
        itens = sorted(secoes[turma], key=lambda x: x[0])
        corrigir = [x for x in itens if x[0] != INFORMACAO]
        info = [x for x in itens if x[0] == INFORMACAO]
        print(f"{turma}\t{len(corrigir)}\t{len(info)}")
        md = [f"# {turma} — o que conferir na planilha de controle", "",
              f"Rodada da sincronização (exportação {rodada.name}). **{len(corrigir)} itens a conferir** e {len(info)} para informação.", "",
              "O sistema já carregou os lançamentos como estão; a correção feita na planilha entra na próxima rodada."
              + (f" Desde {br(virada[turma])} esta turma lança direto no sistema: corrija na planilha só o que é anterior a essa data."
                 if turma in virada else "")]
        grupo = None
        n = 0
        for x in corrigir:
            if x[1] != grupo:
                grupo = x[1]
                md += ["", f"## {grupo}", "", "| # | Prioridade | Onde | O que está errado | O que fazer | Resposta do operador |",
                       "|---:|---|---|---|---|---|"]
            n += 1
            md.append(f"| {n} | {PRIORIDADE[x[0]]} | {x[2]} | {x[3]} | {x[4]} | |")
        if info:
            md += ["", "## Para informação — não corrigir na planilha", ""]
            grupo = None
            for x in info:
                if x[1] != grupo:
                    grupo = x[1]
                    md += ["", f"### {grupo}", ""]
                md.append(f"- {x[2]}: {x[3]} {x[4]}")
        nome = f"{i:02d}-{turma.replace(' ', '-')}"
        (saida / f"{nome}.md").write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")
        folha = livro.create_sheet(turma[:31])
        folha.append(["#", "Prioridade", "Assunto", "Onde", "O que está errado", "O que fazer", "Resposta do operador"])
        for c in folha[1]:
            c.font = Font(bold=True)
        for k, x in enumerate(corrigir + info, start=1):
            folha.append([k, PRIORIDADE[x[0]], x[1], x[2], x[3], x[4], ""])
        for col, largura in zip("ABCDEFG", (5, 13, 28, 34, 70, 60, 30), strict=True):
            folha.column_dimensions[col].width = largura
        for row in folha.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        folha.freeze_panes = "A2"
    livro.save(saida / "erros-por-turma.xlsx")
    sobra = [i["id"] for i in decididos.carregar()["itens"] if i["id"] not in usados]
    print(f"\ndecididos em {decididos.data()}: {len(usados)} casaram com a lista; "
          f"{len(sobra)} nao aparecem mais (sumiram da lista ou mudaram de texto): {', '.join(sobra) or '—'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
